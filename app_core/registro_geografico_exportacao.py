"""Exportação dos RGs selecionados, sem resumos nem metadados técnicos."""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app_core import registro_geografico as rg_core


COLUNAS = (
    "Localidade", "Quarteirão", "Logradouro", "Número", "Sequência", "Lado",
    "Tipo", "Condomínio (unidades)", "Observação", "Data de atualização", "Agentes",
)


def _data(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value:
        try:
            return date.fromisoformat(str(value))
        except ValueError:
            return str(value)  # Preserva datas legadas que não estejam em ISO.
    return None


def gerar_xlsx(db_path, localidade, quarteiroes, base_dir=None):
    """Uma linha por imóvel, na ordem do RG, somente dos quarteirões pedidos."""
    try:
        localidade = int(str(localidade).strip())
        if localidade <= 0:
            raise ValueError
    except (TypeError, ValueError):
        raise ValueError("Informe uma localidade válida.") from None
    selecionados = [str(q).strip() for q in quarteiroes if str(q or "").strip()]
    if not selecionados:
        raise ValueError("Selecione ao menos um quarteirão.")

    disponiveis = {}
    for item in rg_core.quarteiroes_por_localidade(db_path, localidade, base_dir):
        for chave in (item["quarteirao"], item["quarteirao_raw"]):
            disponiveis[chave] = item["quarteirao_raw"]
    canonicos = []
    vistos = set()
    for numero in selecionados:
        if numero not in disponiveis:
            raise ValueError(f"Quarteirão {numero} não encontrado nesta localidade.")
        canonico = disponiveis[numero]
        if canonico not in vistos:
            canonicos.append(canonico)
            vistos.add(canonico)

    wb = Workbook()
    ws = wb.active
    ws.title = "RGs"
    ws.append(COLUNAS)
    linha = 2
    for numero in canonicos:
        dados = rg_core.quarteirao(db_path, localidade, numero, base_dir)
        # Um RG vazio mantém sua identificação, sem inventar um imóvel.
        for registro in dados["registros"] or [{}]:
            valores = (
                dados["localidade"]["nome"], dados["quarteirao"],
                registro.get("logradouro"), registro.get("numero"),
                registro.get("sequencia"), registro.get("lado"), registro.get("tipo"),
                registro.get("condominio"), registro.get("observacao"),
                _data(registro.get("data_atualizacao")), registro.get("agentes"),
            )
            ws.append(valores)
            for coluna in range(1, len(COLUNAS) + 1):
                cell = ws.cell(linha, coluna)
                if isinstance(cell.value, str):
                    # Texto literal: preserva códigos/zeros e nunca executa fórmulas.
                    cell.data_type = "s"
                    cell.number_format = "@"
                elif isinstance(cell.value, date):
                    cell.number_format = "dd/mm/yyyy"
            linha += 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for indice, largura in enumerate((24, 14, 45, 14, 14, 10, 10, 24, 55, 22, 35), 1):
        ws.column_dimensions[get_column_letter(indice)].width = largura
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
