"""XLSForm por ciclo: download local, sem publicação automática no Kobo."""

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill

from app_core import liraa, liraa_kobo, liraa_operacional as op, normalizadores


def _rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:] if any(v is not None for v in row)]


def gerar(target, id_ciclo, base_dir=None):
    plano = op.plano(target, id_ciclo, base_dir)
    config = op.configuracao(target, id_ciclo)
    if not plano["quarteiroes"] or plano["sem_sorteio"]:
        raise liraa.LiraaError("Sorteie todos os estratos do ciclo antes de gerar o formulário.")
    if not config["acs_definidos"]:
        raise liraa.LiraaError("Salve os ACS participantes, mesmo que não haja nenhum neste ciclo.")
    source = load_workbook(op.BASE_XLSFORM, read_only=True)
    try:
        survey, choices = _rows(source["survey"]), _rows(source["choices"])
    finally:
        source.close()
    names = {row["name"]: row for row in survey}
    acs_list = names["ACS"]["type"].split()[1]
    loc_list = names["Localidade"]["type"].split()[1]
    loc_codes = {normalizadores.normalizar_localidade(row["label"]): str(row["name"])
                 for row in choices if row["list_name"] == loc_list}
    used_locs = {q["localidade"] for q in plano["quarteiroes"]}
    if any(normalizadores.normalizar_localidade(loc) not in loc_codes for loc in used_locs):
        raise liraa.LiraaError("Uma localidade sorteada não existe no formulário de referência. Revise o modelo.")
    selected = {a["codigo"]: a["nome"] for a in config["acs"]}
    filtered = []
    for row in choices:
        if row["list_name"] == acs_list:
            if row["name"] != "no-acs" and row["name"] not in selected:
                continue
            row["label"] = selected.get(row["name"], "Nenhum ACS")
        if row["list_name"] == loc_list and normalizadores.normalizar_localidade(row["label"]) not in used_locs:
            continue
        filtered.append(row)
    for q in plano["quarteiroes"]:
        filtered.append({"list_name": "endemias_quarteiroes", "name": liraa_kobo._quarteirao(q["quarteirao"]).lstrip("0"),
                         "label": f"Quarteirão {q['quarteirao_exibicao']}",
                         "localidade": loc_codes[normalizadores.normalizar_localidade(q["localidade"])]})
    # Um identificador municipal pode ocorrer em mais de uma localidade: mantém escolhas únicas por código.
    qnames = [r["name"] for r in filtered if r["list_name"] == "endemias_quarteiroes"]
    if len(qnames) != len(set(qnames)):
        raise liraa.LiraaError("Há numeração municipal repetida entre localidades; revise antes de gerar o formulário.")
    filtered += [{"list_name": "endemias_depositos", "name": code, "label": f"{code} — {label}"}
                 for code, label in op.DEPOSITOS.items()]
    filtered += [{"list_name": "endemias_sim_nao", "name": "sim", "label": "Sim"},
                 {"list_name": "endemias_sim_nao", "name": "nao", "label": "Não"}]
    names["Data"]["required"] = "true"
    bounds = []
    for field, operator in (("inicio", ">="), ("fim", "<=")):
        if plano["ciclo"].get(field):
            bounds.append(f". {operator} date('{plano['ciclo'][field]}')")
    if bounds:
        names["Data"]["constraint"] = " and ".join(bounds)
        names["Data"]["constraint_message"] = "Informe uma data dentro do período deste ciclo."
    names["ACS"]["constraint"] = "not(selected(., 'no-acs') and count-selected(.) > 1)"
    names["ACS"]["constraint_message"] = "Nenhum ACS não pode ser selecionado junto com um nome."
    names["ACS"]["default"] = ""
    names["Localidade"]["default"] = ""
    names["Agentes"]["constraint"] = "not(selected(., 'no-ace') and count-selected(.) > 1)"
    names["Agentes"]["constraint_message"] = "Nenhum ACE não pode ser selecionado junto com um nome."
    names["Quarteir_o"].update(type="select_one endemias_quarteiroes", required="true", default="",
                             choice_filter="localidade=${Localidade}", appearance="minimal",
                             hint="Somente quarteirões sorteados para a localidade escolhida.")
    names["N_mero_do_tubito"].update(required="true", constraint=". > 0",
                                    constraint_message="Informe o número positivo do tubito.")
    names["C_digo_do_dep_sito"].update(type="select_one endemias_depositos", required="true", appearance="minimal")
    names["group_jr1vc40"].update(relevant="${Houve_coleta} = 'sim'", repeat_count="${Quantidade_coletas}")
    names["Morador"]["required"] = "false"
    # Tratamento pertence ao imóvel, não a cada recipiente; a referência antiga a Visita era inválida.
    treatment = []
    start = next(i for i, r in enumerate(survey) if r.get("name") == "group_rb5ho54")
    end = next(i for i in range(start + 1, len(survey)) if survey[i]["type"] == "end_group")
    treatment = survey[start:end + 1]
    treatment[0]["relevant"] = ""
    survey = survey[:start] + survey[end + 1:]
    repeat_index = next(i for i, row in enumerate(survey) if row["type"] == "begin_repeat")
    survey[repeat_index:repeat_index] = [
        {"type": "select_one endemias_sim_nao", "name": "Houve_coleta", "label": "Houve coleta de tubitos?", "required": "true"},
        {"type": "integer", "name": "Quantidade_coletas", "label": "Quantidade de recipientes / tubitos", "required": "true",
         "relevant": "${Houve_coleta} = 'sim'", "constraint": ". >= 1 and . <= 100", "constraint_message": "Informe entre 1 e 100 tubitos."}]
    survey.extend(treatment)
    survey[2:2] = [{"type": "hidden", "name": "endemias_ciclo", "default": str(id_ciclo)},
                   {"type": "hidden", "name": "endemias_plano", "default": plano["plano_hash"]},
                   {"type": "note", "name": "endemias_ciclo_info", "label": f"LIRAa — {plano['ciclo']['ano']} / {plano['ciclo']['nome']}"}]
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in (("survey", survey), ("choices", filtered)):
        sheet = workbook.create_sheet(title)
        headers = list(dict.fromkeys(key for row in rows for key in row if key))
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(key) for key in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="155E75")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(70, max(18, len(str(column[0].value)) + 3))
    settings = workbook.create_sheet("settings")
    settings.append(["form_title", "form_id", "version", "style"])
    settings.append([f"LIRAa {plano['ciclo']['ano']} — {plano['ciclo']['nome']}",
                     f"endemias_liraa_ciclo_{id_ciclo}", datetime.now().strftime("%Y%m%d%H%M%S"), "pages"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
