"""Cadastro de microáreas ACS composto pelos quarteirões da camada ativa."""

import io
import re
from datetime import datetime
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app_core import db as db_core
from app_core import registro_geografico as rg_core
from app_core.excel import excel_safe


class MicroareaError(ValueError):
    pass


def _schema(conn):
    if getattr(conn, "backend", "sqlite") != "sqlite":
        return
    with conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS territorializacao_microareas (
            id_microarea INTEGER PRIMARY KEY AUTOINCREMENT,
            id_localidade INTEGER NOT NULL REFERENCES localidades(id_localidade),
            numero TEXT NOT NULL, acs_codigo TEXT, observacoes TEXT NOT NULL DEFAULT '',
            criado_em TEXT NOT NULL, atualizado_em TEXT NOT NULL,
            UNIQUE(id_localidade, numero))""")
        conn.execute("""CREATE TABLE IF NOT EXISTS territorializacao_microarea_quarteiroes (
            id_microarea INTEGER NOT NULL REFERENCES territorializacao_microareas(id_microarea) ON DELETE CASCADE,
            id_localidade INTEGER NOT NULL REFERENCES localidades(id_localidade),
            quarteirao TEXT NOT NULL,
            PRIMARY KEY(id_microarea, id_localidade, quarteirao),
            UNIQUE(id_localidade, quarteirao))""")
        conn.execute("""CREATE TABLE IF NOT EXISTS acs_catalogo (
            acs_codigo TEXT PRIMARY KEY, nome TEXT NOT NULL, atualizado_em TEXT NOT NULL)""")


def _codigo_quarteirao(valor):
    return rg_core._quarteirao(valor)


def _numero(valor):
    texto = str(valor or "").strip()
    if not re.fullmatch(r"[0-9]+", texto) or int(texto) < 1:
        raise MicroareaError("Informe um número inteiro positivo para a microárea.")
    return str(int(texto))


def _localidade(conn, valor):
    try:
        identificador = int(valor)
    except (ValueError, TypeError):
        raise MicroareaError("Selecione uma localidade válida.") from None
    row = conn.execute("SELECT id_localidade, nome FROM localidades WHERE id_localidade=?", (identificador,)).fetchone()
    if not row:
        raise MicroareaError("Localidade não encontrada.")
    return db_core.serialize_row(row)


def _geometrias(target, base_dir):
    camada = rg_core.geojson_ativo(target, base_dir)
    return {
        (int(f["properties"]["Localidade"]), _codigo_quarteirao(f["properties"]["id_quart"])): f
        for f in camada.get("features", [])
    }


def listar(target, base_dir=None):
    geometrias = _geometrias(target, base_dir)
    conn = db_core.connect(target)
    try:
        _schema(conn)
        rows = conn.execute("""SELECT m.*, l.nome AS localidade, a.nome AS acs_nome
            FROM territorializacao_microareas m
            JOIN localidades l ON l.id_localidade=m.id_localidade
            LEFT JOIN acs_catalogo a ON a.acs_codigo=m.acs_codigo
            ORDER BY l.nome, CAST(m.numero AS INTEGER), m.numero""").fetchall()
        membros = conn.execute("""SELECT id_microarea, id_localidade, quarteirao
            FROM territorializacao_microarea_quarteiroes ORDER BY id_microarea, quarteirao""").fetchall()
        acs = [db_core.serialize_row(row) for row in conn.execute(
            "SELECT acs_codigo, nome FROM acs_catalogo ORDER BY nome, acs_codigo").fetchall()]
        por_id = {}
        for row in membros:
            por_id.setdefault(row["id_microarea"], []).append(row["quarteirao"])
        registros = []
        for row in rows:
            item = db_core.serialize_row(row)
            item["quarteiroes"] = por_id.get(item["id_microarea"], [])
            item["quarteiroes_sem_geometria"] = [q for q in item["quarteiroes"] if (item["id_localidade"], q) not in geometrias]
            registros.append(item)
        return {"registros": registros, "geometrias": len(geometrias), "acs": acs}
    finally:
        conn.close()


def _montar_relatorio(dados, resumo_rg, ids=None, incluir_condominios=True):
    """Agrega o cadastro sem confundir ausência de RG com população zero."""
    todos = dados["registros"]
    ids_selecionados = set(ids) if ids is not None else None
    selecionados = [r for r in todos if ids_selecionados is None or r["id_microarea"] in ids_selecionados]
    rg_quarteiroes = resumo_rg.get("quarteiroes", {})
    acs_catalogo = {a["acs_codigo"]: a["nome"] for a in dados["acs"]}
    acs_vinculados_global = {r["acs_codigo"] for r in todos if r["acs_codigo"]}
    sem_area_global = [
        {"acs_codigo": codigo, "nome": nome}
        for codigo, nome in acs_catalogo.items() if codigo not in acs_vinculados_global
    ]
    registros = []
    por_acs = {}
    por_localidade = {}
    for original in selecionados:
        r = dict(original)
        detalhes = []
        for q in r["quarteiroes"]:
            chave = f"{r['id_localidade']}:{rg_core._quarteirao_display(q)}"
            rg = rg_quarteiroes.get(chave)
            pop_com = int(rg["populacao_aproximada"] or 0) if rg else None
            pop_sem = int(rg["populacao_sem_condominios"] or 0) if rg else None
            detalhes.append({
                "quarteirao": q, "tem_rg": rg is not None,
                "populacao_com_condominios": pop_com,
                "populacao_sem_condominios": pop_sem,
                "populacao_aproximada": pop_com if incluir_condominios else pop_sem,
                "residencias_condominio": int(rg["residencias_condominio"] or 0) if rg else None,
                "imoveis_rg": int(rg["imoveis"] or 0) if rg else None,
            })
        r["detalhes_quarteiroes"] = detalhes
        r["quarteiroes_com_rg"] = sum(d["tem_rg"] for d in detalhes)
        r["quarteiroes_sem_rg"] = len(detalhes) - r["quarteiroes_com_rg"]
        r["populacao_aproximada"] = sum(d["populacao_aproximada"] or 0 for d in detalhes)
        r["populacao_sem_condominios"] = sum(d["populacao_sem_condominios"] or 0 for d in detalhes)
        r["populacao_com_condominios"] = sum(d["populacao_com_condominios"] or 0 for d in detalhes)
        r["residencias_condominio"] = sum(d["residencias_condominio"] or 0 for d in detalhes)
        r["imoveis_rg"] = sum(d["imoveis_rg"] or 0 for d in detalhes)
        registros.append(r)
        local = por_localidade.setdefault(r["id_localidade"], {
            "id_localidade": r["id_localidade"], "localidade": r["localidade"],
            "microareas": 0, "sem_acs": 0, "quarteiroes": 0,
            "quarteiroes_sem_rg": 0, "populacao_aproximada": 0, "acs_codigos": set(),
        })
        local["microareas"] += 1
        local["sem_acs"] += not bool(r["acs_codigo"])
        local["quarteiroes"] += len(detalhes)
        local["quarteiroes_sem_rg"] += r["quarteiroes_sem_rg"]
        local["populacao_aproximada"] += r["populacao_aproximada"]
        if r["acs_codigo"]:
            local["acs_codigos"].add(r["acs_codigo"])
            acs = por_acs.setdefault(r["acs_codigo"], {
                "acs_codigo": r["acs_codigo"],
                "acs_nome": r["acs_nome"] or acs_catalogo.get(r["acs_codigo"]) or r["acs_codigo"],
                "microareas": 0, "quarteiroes": 0, "quarteiroes_com_rg": 0,
                "quarteiroes_sem_rg": 0, "populacao_aproximada": 0,
            })
            acs["microareas"] += 1
            acs["quarteiroes"] += len(detalhes)
            acs["quarteiroes_com_rg"] += r["quarteiroes_com_rg"]
            acs["quarteiroes_sem_rg"] += r["quarteiroes_sem_rg"]
            acs["populacao_aproximada"] += r["populacao_aproximada"]
    acs_rows = sorted(por_acs.values(), key=lambda row: (-row["populacao_aproximada"], row["acs_nome"]))
    locais = []
    for local in sorted(por_localidade.values(), key=lambda row: row["localidade"]):
        local["acs_distintos"] = len(local.pop("acs_codigos"))
        locais.append(local)
    acs_com_rg = [r for r in acs_rows if r["quarteiroes_com_rg"]]
    atribuida = sum(r["populacao_aproximada"] for r in acs_rows)
    indicadores = {
        "microareas": len(registros),
        "com_acs": sum(bool(r["acs_codigo"]) for r in registros),
        "sem_acs": sum(not r["acs_codigo"] for r in registros),
        "acs_distintos": len(acs_rows),
        "acs_catalogo_sem_area_global": len(sem_area_global),
        "diferenca_cadastral_1a1": max(0, sum(not r["acs_codigo"] for r in registros) - len(sem_area_global)),
        "quarteiroes": sum(len(r["quarteiroes"]) for r in registros),
        "quarteiroes_sem_rg": sum(r["quarteiroes_sem_rg"] for r in registros),
        "populacao_aproximada": sum(r["populacao_aproximada"] for r in registros),
        "populacao_atribuida": atribuida,
        "residencias_condominio": sum(r["residencias_condominio"] for r in registros),
        "media_populacao_por_acs_com_rg": round(atribuida / len(acs_com_rg), 1) if acs_com_rg else None,
        "acs_com_rg": len(acs_com_rg),
    }
    return {
        "registros": registros, "acs": dados["acs"], "geometrias": dados["geometrias"],
        "indicadores": indicadores, "por_acs": acs_rows, "por_localidade": locais,
        "acs_sem_area_global": sem_area_global,
        "resumo_rg": resumo_rg,
        "fonte_populacao": resumo_rg.get("fonte_populacao", ""),
        "media_pessoas_por_residencia": resumo_rg.get("media_pessoas_por_residencia"),
        "incluir_condominios": incluir_condominios,
    }


def relatorio(target, base_dir=None, ids=None, incluir_condominios=True):
    dados = listar(target, base_dir)
    resumo_rg = rg_core.resumo_mapa(target, base_dir)
    return _montar_relatorio(dados, resumo_rg, ids, incluir_condominios)


def salvar(target, dados, base_dir=None, id_microarea=None):
    geometrias = _geometrias(target, base_dir)
    conn = db_core.connect(target)
    try:
        _schema(conn)
        localidade = _localidade(conn, dados.get("id_localidade"))
        numero = _numero(dados.get("numero"))
        codigo = str(dados.get("acs_codigo") or "").strip() or None
        observacoes = str(dados.get("observacoes") or "").strip()[:2000]
        if codigo and not conn.execute("SELECT 1 FROM acs_catalogo WHERE acs_codigo=?", (codigo,)).fetchone():
            raise MicroareaError("ACS não encontrado no catálogo local.")
        raw = dados.get("quarteiroes")
        if not isinstance(raw, list) or not raw:
            raise MicroareaError("Selecione ao menos um quarteirão.")
        quarteiroes = [_codigo_quarteirao(q) for q in raw]
        if len(set(quarteiroes)) != len(quarteiroes) or any(not q for q in quarteiroes):
            raise MicroareaError("A seleção contém quarteirões inválidos ou repetidos.")
        atual = None
        if id_microarea is not None:
            atual = conn.execute("SELECT * FROM territorializacao_microareas WHERE id_microarea=?", (id_microarea,)).fetchone()
            if not atual:
                raise MicroareaError("Microárea não encontrada.")
        # Um vínculo preexistente sem geometria pode continuar na edição, mas
        # nunca se pode incluir um novo quarteirão ausente da camada atual.
        antigos = set()
        if atual:
            antigos = {row[0] for row in conn.execute(
                "SELECT quarteirao FROM territorializacao_microarea_quarteiroes WHERE id_microarea=?",
                (id_microarea,),
            )}
        for q in quarteiroes:
            if (localidade["id_localidade"], q) not in geometrias and not (atual and atual["id_localidade"] == localidade["id_localidade"] and q in antigos):
                raise MicroareaError(f"Quarteirão {q} não está na camada geográfica ativa desta localidade.")
        conflito_numero = conn.execute(
            "SELECT id_microarea FROM territorializacao_microareas WHERE id_localidade=? AND numero=? AND id_microarea<>?",
            (localidade["id_localidade"], numero, id_microarea or 0),
        ).fetchone()
        if conflito_numero:
            raise MicroareaError("Já existe uma microárea com esse número na localidade.")
        for q in quarteiroes:
            conflito = conn.execute("""SELECT id_microarea FROM territorializacao_microarea_quarteiroes
                WHERE id_localidade=? AND quarteirao=? AND id_microarea<>?""",
                (localidade["id_localidade"], q, id_microarea or 0),
            ).fetchone()
            if conflito:
                raise MicroareaError(f"Quarteirão {q} já pertence a outra microárea.")
        agora = datetime.now().isoformat(timespec="seconds")
        try:
            with conn:
                if atual:
                    conn.execute("""UPDATE territorializacao_microareas SET id_localidade=?, numero=?,
                        acs_codigo=?, observacoes=?, atualizado_em=? WHERE id_microarea=?""",
                        (localidade["id_localidade"], numero, codigo, observacoes, agora, id_microarea))
                    conn.execute("DELETE FROM territorializacao_microarea_quarteiroes WHERE id_microarea=?", (id_microarea,))
                else:
                    id_microarea = db_core.insert_and_get_id(conn,
                        """INSERT INTO territorializacao_microareas
                        (id_localidade, numero, acs_codigo, observacoes, criado_em, atualizado_em)
                        VALUES (?, ?, ?, ?, ?, ?)""",
                        (localidade["id_localidade"], numero, codigo, observacoes, agora, agora), "id_microarea")
                conn.executemany("""INSERT INTO territorializacao_microarea_quarteiroes
                    (id_microarea, id_localidade, quarteirao) VALUES (?, ?, ?)""",
                    [(id_microarea, localidade["id_localidade"], q) for q in quarteiroes])
        except Exception as exc:
            # As UNIQUE constraints também protegem contra gravações concorrentes.
            if "unique" in str(exc).lower() or "duplicate" in str(exc).lower():
                raise MicroareaError("Número ou quarteirão já utilizado; atualize a lista e tente novamente.") from exc
            raise
        return id_microarea
    finally:
        conn.close()


def excluir(target, id_microarea):
    conn = db_core.connect(target)
    try:
        _schema(conn)
        with conn:
            row = conn.execute("SELECT id_microarea FROM territorializacao_microareas WHERE id_microarea=?", (id_microarea,)).fetchone()
            if not row:
                raise MicroareaError("Microárea não encontrada.")
            conn.execute("DELETE FROM territorializacao_microarea_quarteiroes WHERE id_microarea=?", (id_microarea,))
            conn.execute("DELETE FROM territorializacao_microareas WHERE id_microarea=?", (id_microarea,))
    finally:
        conn.close()


def _selecionadas(target, base_dir, ids=None):
    registros = listar(target, base_dir)["registros"]
    if ids is not None:
        registros = [r for r in registros if r["id_microarea"] in ids]
    return registros


def exportar_xlsx(target, base_dir=None, ids=None, incluir_condominios=True):
    dados = relatorio(target, base_dir, ids, incluir_condominios)
    registros = dados["registros"]
    indicadores = dados["indicadores"]
    wb = Workbook()
    painel = wb.active
    painel.title = "Indicadores"
    painel.append(["Indicador", "Valor"])
    painel.append(["População de condomínios residenciais", "Incluída" if incluir_condominios else "Excluída"])
    for nome, valor in [
        ("Microáreas no recorte", indicadores["microareas"]),
        ("Microáreas com ACS", indicadores["com_acs"]),
        ("Microáreas sem ACS", indicadores["sem_acs"]),
        ("ACS distintos no recorte", indicadores["acs_distintos"]),
        ("ACS do catálogo sem microárea (global)", indicadores["acs_catalogo_sem_area_global"]),
        ("Diferença cadastral 1:1 (não é déficit de pessoal)", indicadores["diferenca_cadastral_1a1"]),
        ("Quarteirões no recorte", indicadores["quarteiroes"]),
        ("Quarteirões sem RG", indicadores["quarteiroes_sem_rg"]),
        ("Unidades residenciais em condomínios", indicadores["residencias_condominio"]),
        ("População estimada nas microáreas com RG", indicadores["populacao_aproximada"]),
        ("População estimada atribuída a ACS", indicadores["populacao_atribuida"]),
        ("Média estimada por ACS com RG", indicadores["media_populacao_por_acs_com_rg"]),
        ("ACS no cálculo da média", indicadores["acs_com_rg"]),
    ]:
        painel.append([nome, valor])
    painel.append(["Fonte", f"RG: residências × {dados['media_pessoas_por_residencia']} pessoas; {dados['fonte_populacao']}"])
    painel.append(["Limite", "Quarteirões sem RG não entram na estimativa; ACS sem microárea no catálogo não significam disponibilidade de pessoal. Condomínios são imóveis residenciais com condominio > 0 no RG."])
    por_acs = wb.create_sheet("Por ACS")
    por_acs.append(["ACS código", "ACS nome", "Microáreas", "Quarteirões", "Quarteirões sem RG", "População estimada"])
    for item in dados["por_acs"]:
        por_acs.append([excel_safe(item["acs_codigo"]), excel_safe(item["acs_nome"]), item["microareas"],
                        item["quarteiroes"], item["quarteiroes_sem_rg"],
                        item["populacao_aproximada"] if item["quarteiroes_com_rg"] else None])
    por_localidade = wb.create_sheet("Por localidade")
    por_localidade.append(["Localidade", "Microáreas", "Sem ACS", "ACS distintos", "Quarteirões", "Quarteirões sem RG", "População estimada"])
    for item in dados["por_localidade"]:
        por_localidade.append([excel_safe(item["localidade"]), item["microareas"], item["sem_acs"],
                               item["acs_distintos"], item["quarteiroes"], item["quarteiroes_sem_rg"],
                               item["populacao_aproximada"] if item["quarteiroes"] > item["quarteiroes_sem_rg"] else None])
    resumo = wb.create_sheet("Microáreas")
    resumo.append(["Localidade", "Microárea", "ACS código", "ACS nome", "Quarteirões", "Sem geometria", "Observações", "Atualizado em",
                   "Quarteirões sem RG", "População estimada", "Imóveis RG"])
    membros = wb.create_sheet("Quarteirões")
    membros.append(["Localidade", "Microárea", "Quarteirão", "ACS código", "ACS nome", "Sem geometria", "RG cadastrado", "População estimada", "Imóveis RG"])
    for r in registros:
        resumo.append([excel_safe(r["localidade"]), excel_safe(r["numero"]), excel_safe(r["acs_codigo"]),
                       excel_safe(r["acs_nome"]), len(r["quarteiroes"]), len(r["quarteiroes_sem_geometria"]),
                       excel_safe(r["observacoes"]), excel_safe(r["atualizado_em"]), r["quarteiroes_sem_rg"],
                       r["populacao_aproximada"] if r["quarteiroes_com_rg"] else None, r["imoveis_rg"] if r["quarteiroes_com_rg"] else None])
        for detalhe in r["detalhes_quarteiroes"]:
            q = detalhe["quarteirao"]
            membros.append([excel_safe(r["localidade"]), excel_safe(r["numero"]), excel_safe(q),
                            excel_safe(r["acs_codigo"]), excel_safe(r["acs_nome"]), q in r["quarteiroes_sem_geometria"],
                            detalhe["tem_rg"], detalhe["populacao_aproximada"], detalhe["imoveis_rg"]])
    for sheet in wb:
        if sheet.title != "Indicadores":
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 25
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="155E75")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in sheet.columns:
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(55, max(15, max(len(str(cell.value or "")) for cell in column) + 2))
    painel.column_dimensions["A"].width = 52
    painel.column_dimensions["B"].width = 100
    for row in (painel.max_row - 1, painel.max_row):
        painel.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="center")
        painel.row_dimensions[row].height = 31
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def exportar_geojson(target, base_dir=None, ids=None):
    geometrias = _geometrias(target, base_dir)
    features = []
    for r in _selecionadas(target, base_dir, ids):
        for q in r["quarteiroes"]:
            feature = geometrias.get((r["id_localidade"], q))
            if not feature:
                continue
            features.append({"type": "Feature", "geometry": feature["geometry"], "properties": {
                "id_microarea": r["id_microarea"], "microarea": r["numero"],
                "Localidade": r["localidade"], "id_Q": q,
                "acs_codigo": r["acs_codigo"] or "", "acs_nome": r["acs_nome"] or "",
                "observacoes": r["observacoes"],
            }})
    return {"type": "FeatureCollection", "features": features}


def exportar_kml(target, base_dir=None, ids=None):
    geometrias = _geometrias(target, base_dir)
    root = ET.Element("kml", xmlns="http://www.opengis.net/kml/2.2")
    document = ET.SubElement(root, "Document")
    ET.SubElement(document, "name").text = "Microáreas ACS"
    for r in _selecionadas(target, base_dir, ids):
        presentes = [(q, geometrias[(r["id_localidade"], q)]) for q in r["quarteiroes"]
                     if (r["id_localidade"], q) in geometrias]
        if not presentes:
            continue
        placemark = ET.SubElement(document, "Placemark")
        ET.SubElement(placemark, "name").text = f"{r['localidade']} · Microárea {r['numero']}"
        ET.SubElement(placemark, "description").text = f"ACS: {r['acs_nome'] or 'Não atribuído'} | Quarteirões: {', '.join(r['quarteiroes'])} | {r['observacoes']}"
        extras = ET.SubElement(placemark, "ExtendedData")
        for chave, valor in {
            "id_microarea": r["id_microarea"], "localidade": r["localidade"], "microarea": r["numero"],
            "acs_codigo": r["acs_codigo"] or "", "acs_nome": r["acs_nome"] or "",
            "quarteiroes": ", ".join(r["quarteiroes"]), "observacoes": r["observacoes"],
        }.items():
            atributo = ET.SubElement(extras, "Data", name=chave)
            ET.SubElement(atributo, "value").text = str(valor)
        multi = ET.SubElement(placemark, "MultiGeometry")
        for q, feature in presentes:
            geometry = feature["geometry"]
            polygons = [geometry["coordinates"]] if geometry["type"] == "Polygon" else geometry["coordinates"]
            for rings in polygons:
                polygon = ET.SubElement(multi, "Polygon")
                for index, ring in enumerate(rings):
                    boundary = ET.SubElement(polygon, "outerBoundaryIs" if index == 0 else "innerBoundaryIs")
                    line = ET.SubElement(boundary, "LinearRing")
                    ET.SubElement(line, "coordinates").text = " ".join(f"{point[0]},{point[1]}" for point in ring)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)
