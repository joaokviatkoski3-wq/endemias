"""Exportacao analitica do monitoramento de ovitrampas para XLSX."""

import io
from collections import defaultdict
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from app_core import contaovos_registro
from app_core import db as db_core
from app_core import ovitrampas
from app_core.excel import excel_safe


AZUL = "1A4FBA"
AZUL_CLARO = "DCE6F8"
VERDE_CLARO = "DCFCE7"
CINZA_CLARO = "E5E7EB"


LEITURAS_COLUNAS = (
    ("id_contagem", "ID da contagem"),
    ("ovitrampa_id", "ID da ovitrampa"),
    ("ano", "Ano"),
    ("semana", "Semana"),
    ("data", "Data da coleta"),
    ("data_envio_contagem", "Envio ao Conta Ovos"),
    ("ovos", "Ovos"),
    ("leitura_positiva", "Leitura positiva"),
    ("resultado", "Resultado Conta Ovos"),
    ("codigo_conta_ovos", "Código Conta Ovos"),
    ("observacao_conta_ovos", "Observação Conta Ovos"),
    ("api_latitude", "Latitude da leitura (API)"),
    ("api_longitude", "Longitude da leitura (API)"),
    ("api_lat_lng", "Coordenadas da leitura (API)"),
    ("api_arquivo_origem", "Origem da leitura"),
    ("api_importado_em", "Leitura sincronizada em"),
    ("localidade", "Localidade (cadastro local)"),
    ("rua", "Rua (cadastro local)"),
    ("numero", "Numero (cadastro local)"),
    ("complemento", "Complemento (cadastro local)"),
    ("bairro", "Bairro (cadastro local)"),
    ("localizacao", "Localização da ovitrampa"),
    ("quarteirao", "Quarteirão (cadastro local)"),
    ("responsavel", "Responsável"),
    ("telefone_responsavel", "Telefone do responsável"),
    ("local_latitude", "Latitude (cadastro local)"),
    ("local_longitude", "Longitude (cadastro local)"),
    ("cadastro_local_ativo", "Cadastro local ativo"),
    ("cadastro_local_origem", "Origem do cadastro local"),
    ("cadastro_local_atualizado_em", "Cadastro local atualizado em"),
    ("remoto_ovitrap_id", "ID interno remoto"),
    ("remoto_municipio", "Município (cadastro remoto)"),
    ("remoto_municipio_codigo", "Código do município (remoto)"),
    ("remoto_estado", "Estado (cadastro remoto)"),
    ("remoto_latitude", "Latitude (cadastro remoto)"),
    ("remoto_longitude", "Longitude (cadastro remoto)"),
    ("remoto_coordenada_erro", "Erro de coordenada (remoto)"),
    ("remoto_ovos_media", "Média de ovos (cadastro remoto)"),
    ("remoto_quarteirao_id", "ID do quarteirão remoto"),
    ("remoto_grupo_id", "ID do grupo remoto"),
    ("remoto_usuario_id", "ID do usuario remoto"),
    ("remoto_atualizado_em", "Cadastro remoto atualizado em"),
    ("remoto_sincronizado_em", "Cadastro remoto sincronizado em"),
    ("lote_id", "ID do lote de laboratório"),
    ("diario_id", "ID do diário"),
    ("diario_nome", "Diário"),
    ("movimento", "Movimento"),
    ("ciclo", "Ciclo"),
    ("lote_status", "Status do lote"),
    ("laboratorista_id", "ID do laboratorista"),
    ("laboratorista_nome", "Laboratorista"),
    ("laboratorio_ovos", "Ovos lançados no laboratório"),
    ("ocorrencia_codigo", "Código da ocorrência"),
    ("ocorrencia", "Ocorrência"),
    ("lote_iniciado_em", "Lote iniciado em"),
    ("lote_concluido_em", "Lote concluido em"),
    ("lote_enviado_em", "Lote enviado ao Conta Ovos em"),
    ("lote_enviado_por", "Lote enviado por"),
)


ARMADILHAS_COLUNAS = (
    ("ovitrampa_id", "ID da ovitrampa"),
    ("origem_cadastro", "Presença nos cadastros"),
    ("localidade", "Localidade"),
    ("rua", "Rua"),
    ("numero", "Numero"),
    ("complemento", "Complemento"),
    ("bairro", "Bairro"),
    ("localizacao", "Localização da ovitrampa"),
    ("quarteirao", "Quarteirão"),
    ("responsavel", "Responsável"),
    ("telefone_responsavel", "Telefone do responsável"),
    ("local_latitude", "Latitude (cadastro local)"),
    ("local_longitude", "Longitude (cadastro local)"),
    ("cadastro_local_ativo", "Cadastro local ativo"),
    ("cadastro_local_origem", "Origem do cadastro local"),
    ("cadastro_local_atualizado_em", "Cadastro local atualizado em"),
    ("remoto_ovitrap_id", "ID interno remoto"),
    ("remoto_municipio", "Município (cadastro remoto)"),
    ("remoto_municipio_codigo", "Código do município (remoto)"),
    ("remoto_estado", "Estado (cadastro remoto)"),
    ("remoto_latitude", "Latitude (cadastro remoto)"),
    ("remoto_longitude", "Longitude (cadastro remoto)"),
    ("remoto_coordenada_erro", "Erro de coordenada (remoto)"),
    ("remoto_ovos_media", "Média de ovos (cadastro remoto)"),
    ("remoto_quarteirao_id", "ID do quarteirão remoto"),
    ("remoto_grupo_id", "ID do grupo remoto"),
    ("remoto_usuario_id", "ID do usuario remoto"),
    ("remoto_atualizado_em", "Cadastro remoto atualizado em"),
    ("remoto_sincronizado_em", "Cadastro remoto sincronizado em"),
    ("leituras", "Leituras no período"),
    ("positivas", "Leituras positivas no período"),
    ("ovos", "Ovos no período"),
    ("ipo", "IPO (%)"),
    ("ido", "IDO"),
    ("imo", "IMO"),
    ("primeira_leitura", "Primeira leitura no período"),
    ("ultima_leitura", "Última leitura no período"),
    ("ultima_positiva", "Última leitura positiva no período"),
    ("atende_ranking", "Atende aos filtros do ranking"),
)


def _dict_rows(cursor):
    return [db_core.serialize_row(row) for row in cursor.fetchall()]


def _lab_sql(conn):
    if not (
        db_core.table_exists(conn, ovitrampas.LAB_LOTES_TABLE)
        and db_core.table_exists(conn, ovitrampas.LAB_ITENS_TABLE)
    ):
        campos = ", ".join(
            f"NULL AS {nome}" for nome in (
                "lote_id", "diario_id", "diario_nome", "movimento", "ciclo",
                "lote_status", "laboratorista_id", "laboratorista_nome",
                "laboratorio_ovos", "ocorrencia_codigo", "lote_iniciado_em",
                "lote_concluido_em", "lote_enviado_em", "lote_enviado_por",
            )
        )
        return "", "", campos

    cte = f"""
        lab AS (
            SELECT lt.id_lote AS lote_id, lt.id_diario AS diario_id,
                   lt.diario_nome, lt.data_movimento, lt.movimento, lt.ciclo,
                   lt.status AS lote_status,
                   lt.id_laboratorista AS laboratorista_id,
                   lt.laboratorista_nome, li.ovitrampa_id,
                   li.ovos AS laboratorio_ovos,
                   li.ocorrencia AS ocorrencia_codigo,
                   lt.iniciado_em AS lote_iniciado_em,
                   lt.concluido_em AS lote_concluido_em,
                   lt.enviado_conta_ovos_em AS lote_enviado_em,
                   lt.enviado_por_nome AS lote_enviado_por,
                   ROW_NUMBER() OVER (
                       PARTITION BY li.ovitrampa_id, lt.data_movimento
                       ORDER BY CASE lt.status
                           WHEN 'enviado_conta_ovos' THEN 3
                           WHEN 'concluido' THEN 2 ELSE 1 END DESC,
                           lt.atualizado_em DESC, lt.id_lote DESC
                   ) AS rn
              FROM {ovitrampas.LAB_ITENS_TABLE} li
              JOIN {ovitrampas.LAB_LOTES_TABLE} lt ON lt.id_lote=li.id_lote
             WHERE lt.status IN ('concluido','enviado_conta_ovos')
        )
    """
    join = """
        LEFT JOIN lab
          ON lab.ovitrampa_id=o.ovitrampa_id AND lab.data_movimento=o.data
         AND lab.rn=1
    """
    campos = ", ".join(
        f"lab.{nome}" for nome in (
            "lote_id", "diario_id", "diario_nome", "movimento", "ciclo",
            "lote_status", "laboratorista_id", "laboratorista_nome",
            "laboratorio_ovos", "ocorrencia_codigo", "lote_iniciado_em",
            "lote_concluido_em", "lote_enviado_em", "lote_enviado_por",
        )
    )
    return cte, join, campos


def _leituras(conn, preparados):
    contaovos_registro.ensure_schema_connection(conn)
    lab_cte, lab_join, lab_campos = _lab_sql(conn)
    with_sql = f"WITH {lab_cte}" if lab_cte else ""
    rows = _dict_rows(conn.execute(
        f"""{with_sql}
        SELECT o.id_contagem, o.ovitrampa_id, o.ano, o.semana, o.data,
               o.data_envio_contagem, o.ovos,
               CASE WHEN COALESCE(o.ovos,0)>0 THEN 'Sim' ELSE 'Não' END AS leitura_positiva,
               o.resultado, o.codigo_conta_ovos, o.observacao_conta_ovos,
               o.latitude AS api_latitude, o.longitude AS api_longitude,
               o.lat_lng AS api_lat_lng, o.arquivo_origem AS api_arquivo_origem,
               o.importado_em AS api_importado_em,
               am.localidade, am.rua, am.numero, am.complemento, am.bairro,
               am.localizacao, am.quarteirao, am.responsavel,
               am.telefone_responsavel, am.latitude AS local_latitude,
               am.longitude AS local_longitude, am.ativo AS cadastro_local_ativo,
               am.arquivo_origem AS cadastro_local_origem,
               am.atualizado_em AS cadastro_local_atualizado_em,
               r.ovitrap_id AS remoto_ovitrap_id, r.municipio AS remoto_municipio,
               r.municipio_codigo AS remoto_municipio_codigo,
               r.estado AS remoto_estado, r.latitude AS remoto_latitude,
               r.longitude AS remoto_longitude,
               r.coordenada_erro AS remoto_coordenada_erro,
               r.ovos_media AS remoto_ovos_media,
               r.quarteirao_remoto_id AS remoto_quarteirao_id,
               r.grupo_remoto_id AS remoto_grupo_id,
               r.usuario_remoto_id AS remoto_usuario_id,
               r.atualizado_remoto_em AS remoto_atualizado_em,
               r.sincronizado_em AS remoto_sincronizado_em,
               {lab_campos}
          FROM {ovitrampas.OCORRENCIAS_TABLE} o
          LEFT JOIN {ovitrampas.ARMADILHAS_TABLE} am
            ON am.ovitrampa_id=o.ovitrampa_id
          LEFT JOIN {contaovos_registro.TABLE} r
            ON r.ovitrampa_id_remoto=o.ovitrampa_id
          {lab_join}
          {preparados['where']}
         ORDER BY o.ano DESC, o.semana DESC, o.data DESC, o.ovitrampa_id""",
        preparados["params"],
    ))
    for row in rows:
        codigo = row.get("ocorrencia_codigo")
        row["ocorrencia"] = ovitrampas.OCORRENCIAS.get(codigo) if codigo else None
    return rows


def _armadilhas(conn, preparados):
    contaovos_registro.ensure_schema_connection(conn)
    distrito = preparados["distrito"]
    termo = preparados["ovitrampa"]
    outer_clauses = ["1=1"]
    outer_params = []
    if distrito:
        outer_clauses.append("am.localidade=?")
        outer_params.append(distrito)
    if termo:
        outer_clauses.append("LOWER(ids.ovitrampa_id) LIKE ?")
        outer_params.append(f"%{termo.lower()}%")
    positivas = "SUM(CASE WHEN COALESCE(o.ovos,0)>0 THEN 1 ELSE 0 END)"
    ovos = "COALESCE(SUM(o.ovos),0)"
    ipo = ovitrampas._round_one(
        f"CASE WHEN COUNT(*)>0 THEN 100.0*{positivas}/COUNT(*) ELSE 0 END"
    )
    ido = ovitrampas._round_one(
        f"CASE WHEN {positivas}>0 THEN 1.0*{ovos}/{positivas} ELSE 0 END"
    )
    imo = ovitrampas._round_two(
        f"CASE WHEN COUNT(*)>0 THEN 1.0*{ovos}/COUNT(*) ELSE 0 END"
    )
    rows = _dict_rows(conn.execute(
        f"""WITH ids AS (
                SELECT ovitrampa_id FROM {ovitrampas.ARMADILHAS_TABLE}
                UNION SELECT ovitrampa_id_remoto FROM {contaovos_registro.TABLE}
                UNION SELECT ovitrampa_id FROM {ovitrampas.OCORRENCIAS_TABLE}
             ), leitura AS (
                SELECT o.ovitrampa_id, COUNT(*) AS leituras,
                       {positivas} AS positivas, {ovos} AS ovos,
                       {ipo} AS ipo, {ido} AS ido, {imo} AS imo,
                       MIN(o.data) AS primeira_leitura,
                       MAX(o.data) AS ultima_leitura,
                       MAX(CASE WHEN COALESCE(o.ovos,0)>0 THEN o.data END) AS ultima_positiva
                  FROM {ovitrampas.OCORRENCIAS_TABLE} o
                  LEFT JOIN {ovitrampas.ARMADILHAS_TABLE} am
                    ON am.ovitrampa_id=o.ovitrampa_id
                  {preparados['where']}
                 GROUP BY o.ovitrampa_id
             )
        SELECT ids.ovitrampa_id,
               CASE WHEN am.ovitrampa_id IS NOT NULL AND r.ovitrampa_id_remoto IS NOT NULL
                    THEN 'Local e Conta Ovos'
                    WHEN am.ovitrampa_id IS NOT NULL THEN 'Somente local'
                    WHEN r.ovitrampa_id_remoto IS NOT NULL THEN 'Somente Conta Ovos'
                    ELSE 'Somente em leituras' END AS origem_cadastro,
               am.localidade, am.rua, am.numero, am.complemento, am.bairro,
               am.localizacao, am.quarteirao, am.responsavel,
               am.telefone_responsavel, am.latitude AS local_latitude,
               am.longitude AS local_longitude, am.ativo AS cadastro_local_ativo,
               am.arquivo_origem AS cadastro_local_origem,
               am.atualizado_em AS cadastro_local_atualizado_em,
               r.ovitrap_id AS remoto_ovitrap_id, r.municipio AS remoto_municipio,
               r.municipio_codigo AS remoto_municipio_codigo,
               r.estado AS remoto_estado, r.latitude AS remoto_latitude,
               r.longitude AS remoto_longitude,
               r.coordenada_erro AS remoto_coordenada_erro,
               r.ovos_media AS remoto_ovos_media,
               r.quarteirao_remoto_id AS remoto_quarteirao_id,
               r.grupo_remoto_id AS remoto_grupo_id,
               r.usuario_remoto_id AS remoto_usuario_id,
               r.atualizado_remoto_em AS remoto_atualizado_em,
               r.sincronizado_em AS remoto_sincronizado_em,
               COALESCE(leitura.leituras,0) AS leituras,
               COALESCE(leitura.positivas,0) AS positivas,
               COALESCE(leitura.ovos,0) AS ovos,
               COALESCE(leitura.ipo,0) AS ipo,
               COALESCE(leitura.ido,0) AS ido,
               COALESCE(leitura.imo,0) AS imo,
               leitura.primeira_leitura, leitura.ultima_leitura,
               leitura.ultima_positiva
          FROM ids
          LEFT JOIN {ovitrampas.ARMADILHAS_TABLE} am
            ON am.ovitrampa_id=ids.ovitrampa_id
          LEFT JOIN {contaovos_registro.TABLE} r
            ON r.ovitrampa_id_remoto=ids.ovitrampa_id
          LEFT JOIN leitura ON leitura.ovitrampa_id=ids.ovitrampa_id
         WHERE {' AND '.join(outer_clauses)}
         ORDER BY ids.ovitrampa_id""",
        [*preparados["params"], *outer_params],
    ))
    for row in rows:
        row["atende_ranking"] = "Sim" if _atende_ranking(row, preparados) else "Não"
    ordenar = preparados["ordenar"]
    ordem = {
        "positivas": ("positivas", "ovos", "ipo"),
        "ipo": ("ipo", "positivas", "ovos"),
        "ido": ("ido", "ovos", "positivas"),
        "ovos": ("ovos", "positivas", "ipo"),
        "leituras": ("leituras", "positivas", "ovos"),
        "recente": ("ultima_positiva", "positivas", "ovos"),
    }.get(ordenar, ("positivas", "ovos", "ipo"))
    rows.sort(
        key=lambda row: tuple(_sort_value(row.get(campo)) for campo in ordem),
        reverse=True,
    )
    return rows


def _ocorrencias_laboratorio(conn, preparados):
    if not (
        db_core.table_exists(conn, ovitrampas.LAB_LOTES_TABLE)
        and db_core.table_exists(conn, ovitrampas.LAB_ITENS_TABLE)
    ):
        return []
    rows = _dict_rows(conn.execute(
        f"""SELECT o.id_contagem, li.ovitrampa_id, o.ano, o.semana, o.data,
                   COALESCE(o.ovos,li.ovos,0) AS ovos,
                   am.localidade, am.rua, am.numero, am.complemento, am.quarteirao,
                   lt.id_lote AS lote_id, lt.diario_nome, lt.movimento, lt.ciclo,
                   lt.status AS lote_status, lt.laboratorista_nome,
                   li.ocorrencia AS ocorrencia_codigo
              FROM {ovitrampas.LAB_ITENS_TABLE} li
              JOIN {ovitrampas.LAB_LOTES_TABLE} lt ON lt.id_lote=li.id_lote
                   AND lt.status IN ('concluido','enviado_conta_ovos')
              JOIN {ovitrampas.OCORRENCIAS_TABLE} o
                ON o.ovitrampa_id=li.ovitrampa_id AND o.data=lt.data_movimento
              LEFT JOIN {ovitrampas.ARMADILHAS_TABLE} am
                ON am.ovitrampa_id=li.ovitrampa_id
              {preparados['where']}
                AND li.ocorrencia IS NOT NULL
             ORDER BY o.ano DESC, o.semana DESC, o.data DESC,
                      li.ovitrampa_id, lt.id_lote""",
        preparados["params"],
    ))
    for row in rows:
        row["ocorrencia"] = ovitrampas.OCORRENCIAS.get(row.get("ocorrencia_codigo"))
    return rows


def _sort_value(value):
    if value is None:
        return (0, "")
    return (1, value)


def _atende_ranking(row, preparados):
    leituras = int(row.get("leituras") or 0)
    positivas = int(row.get("positivas") or 0)
    if positivas <= 0 or leituras < preparados["min_leituras"]:
        return False
    limites = (
        ("ipo", preparados["min_ipo"]),
        ("ido", preparados["min_ido"]),
        ("imo", preparados["min_imo"]),
    )
    return all(limite is None or float(row.get(campo) or 0) >= limite for campo, limite in limites)


def _agrupar_leituras(rows, chave):
    grupos = defaultdict(lambda: {"ids": set(), "leituras": 0, "positivas": 0, "ovos": 0})
    for row in rows:
        valor = chave(row)
        grupo = grupos[valor]
        grupo["ids"].add(row.get("ovitrampa_id"))
        grupo["leituras"] += 1
        if int(row.get("ovos") or 0) > 0:
            grupo["positivas"] += 1
        grupo["ovos"] += int(row.get("ovos") or 0)
    saida = []
    for valor, grupo in grupos.items():
        leituras = grupo["leituras"]
        positivas = grupo["positivas"]
        ovos = grupo["ovos"]
        saida.append({
            "grupo": valor,
            "leituras": leituras,
            "armadilhas": len(grupo["ids"]),
            "positivas": positivas,
            "ovos": ovos,
            "ipo": round(100.0 * positivas / leituras, 1) if leituras else 0,
            "ido": round(ovos / positivas, 1) if positivas else 0,
            "imo": round(ovos / leituras, 2) if leituras else 0,
        })
    return saida


def _excel_value(value):
    if value is None:
        return None
    if isinstance(value, (int, float, date, datetime)):
        return value
    text = str(value).strip()
    if not text:
        return None
    for formato in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            parsed = datetime.strptime(text[:19], formato)
            return parsed.date() if formato == "%Y-%m-%d" else parsed
        except ValueError:
            continue
    return excel_safe(text)


def _write_table_sheet(wb, titulo, colunas, rows, freeze="A2"):
    ws = wb.create_sheet(titulo)
    ws.sheet_view.showGridLines = False
    headers = [label for _, label in colunas]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_excel_value(row.get(key)) for key, _ in colunas])
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for column in ws.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value or "")) for cell in column[:500]) + 2
        ws.column_dimensions[letter].width = min(max(width, 11), 42)
        for cell in column[1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"
    return ws


def _write_resumo(wb, monitor, preparados, armadilhas, gerado_em):
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    ws["A2"] = "Monitoramento de ovitrampas"
    ws["A2"].font = Font(bold=True, size=15, color=AZUL)
    ws["A3"] = "Recorte exportado dos espelhos locais do Conta Ovos e dos dados do sistema"
    ws["A3"].font = Font(italic=True, color="475569")
    ws["A5"] = "Gerado em"
    ws["B5"] = gerado_em
    ws["B5"].number_format = "dd/mm/yyyy hh:mm"

    filtros = preparados["periodo"]
    linhas_filtros = (
        ("Ano", filtros.get("ano")),
        ("Semana inicial", filtros.get("semana_ini")),
        ("Semana final", filtros.get("semana_fim")),
        ("Data inicial", filtros.get("data_ini")),
        ("Data final", filtros.get("data_fim")),
        ("Localidade", preparados.get("distrito") or "Todas"),
        ("Ovitrampa", preparados.get("ovitrampa") or "Todas"),
        ("Mínimo de leituras no ranking", preparados.get("min_leituras")),
        ("IPO mínimo no ranking (%)", preparados.get("min_ipo")),
        ("IDO mínimo no ranking", preparados.get("min_ido")),
        ("IMO mínimo no ranking", preparados.get("min_imo")),
        ("Ordenação do ranking", preparados.get("ordenar")),
    )
    ws["A7"] = "Filtros aplicados"
    ws["A7"].font = Font(bold=True, color="FFFFFF")
    ws["A7"].fill = PatternFill("solid", fgColor=AZUL)
    ws["B7"].fill = PatternFill("solid", fgColor=AZUL)
    for idx, (nome, valor) in enumerate(linhas_filtros, 8):
        ws.cell(idx, 1, nome)
        ws.cell(idx, 2, _excel_value(valor))

    totais = monitor["totais"]
    indicadores = (
        ("Leituras", totais.get("leituras", 0)),
        ("Armadilhas lidas", totais.get("armadilhas_lidas", 0)),
        ("Leituras positivas", totais.get("positivas", 0)),
        ("Armadilhas com leitura positiva", totais.get("armadilhas_positivas", 0)),
        ("Ovos", totais.get("ovos", 0)),
        ("IPO (%)", totais.get("ipo", 0)),
        ("IDO", totais.get("ido", 0)),
        ("IMO", totais.get("imo", totais.get("idv", 0))),
        ("Ocorrências laboratoriais", totais.get("ocorrencias", 0)),
        ("Pendentes de realocação", totais.get("realocar", 0)),
        ("Armadilhas cadastradas no recorte", len(armadilhas)),
        ("Cadastros local e remoto", sum(r["origem_cadastro"] == "Local e Conta Ovos" for r in armadilhas)),
        ("Somente no cadastro local", sum(r["origem_cadastro"] == "Somente local" for r in armadilhas)),
        ("Somente no cadastro Conta Ovos", sum(r["origem_cadastro"] == "Somente Conta Ovos" for r in armadilhas)),
        ("Somente no histórico de leituras", sum(r["origem_cadastro"] == "Somente em leituras" for r in armadilhas)),
    )
    ws["D7"] = "Indicadores do período"
    ws["D7"].font = Font(bold=True, color="FFFFFF")
    ws["D7"].fill = PatternFill("solid", fgColor=AZUL)
    ws["E7"].fill = PatternFill("solid", fgColor=AZUL)
    for idx, (nome, valor) in enumerate(indicadores, 8):
        ws.cell(idx, 4, nome)
        ws.cell(idx, 5, _excel_value(valor))

    fonte_row = max(22, 9 + len(indicadores))
    ws.cell(fonte_row, 1, "Fontes e regras")
    ws.cell(fonte_row, 1).font = Font(bold=True, color="FFFFFF")
    ws.cell(fonte_row, 1).fill = PatternFill("solid", fgColor=AZUL)
    ws.cell(fonte_row, 2).fill = PatternFill("solid", fgColor=AZUL)
    notas = (
        "Leituras e resultados: espelho GET do Conta Ovos, sem consulta remota durante a exportação.",
        "Endereço, responsável, telefone e situação: cadastro local de armadilhas.",
        "Identificadores e coordenadas remotas: espelho do cadastro público do Conta Ovos.",
        "Laboratorista e ocorrências: lotes concluídos ou enviados da aba Laboratório.",
        "IPO = positivas / leituras x 100; IDO = ovos / positivas; IMO = ovos / leituras.",
        "Os limites mínimos de IPO, IDO, IMO e leituras marcam a coluna de ranking; não removem dados brutos.",
    )
    for offset, nota in enumerate(notas, 1):
        ws.cell(fonte_row + offset, 1, nota)
        ws.merge_cells(start_row=fonte_row + offset, start_column=1, end_row=fonte_row + offset, end_column=5)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 3
    ws.column_dimensions["D"].width = 38
    ws.column_dimensions["E"].width = 18
    for row in ws.iter_rows(min_row=8, max_row=fonte_row - 1, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
    return ws


def gerar_monitoramento_xlsx(db_path, filtros=None):
    """Gera um snapshot XLSX completo conforme os filtros do Monitoramento."""
    filtros = filtros or {}
    monitor = ovitrampas.monitoramento_contagens(db_path, filtros)
    conn = db_core.connect(db_path)
    try:
        ovitrampas.ensure_schema(conn)
        preparados = ovitrampas.filtros_monitoramento_contagens(conn, filtros)
        leituras = _leituras(conn, preparados)
        armadilhas = _armadilhas(conn, preparados)
        ocorrencias = _ocorrencias_laboratorio(conn, preparados)
    finally:
        conn.close()

    semanas = _agrupar_leituras(
        leituras,
        lambda row: f"{row.get('ano')} / {int(row.get('semana') or 0):02d}",
    )
    semanas.sort(key=lambda row: row["grupo"], reverse=True)
    localidades = _agrupar_leituras(
        leituras, lambda row: row.get("localidade") or "Sem localidade"
    )
    localidades.sort(key=lambda row: (row["ipo"], row["ido"], row["ovos"]), reverse=True)
    wb = openpyxl.Workbook()
    gerado_em = datetime.now().replace(microsecond=0)
    _write_resumo(wb, monitor, preparados, armadilhas, gerado_em)
    _write_table_sheet(wb, "Leituras", LEITURAS_COLUNAS, leituras, freeze="C2")
    _write_table_sheet(wb, "Armadilhas", ARMADILHAS_COLUNAS, armadilhas, freeze="C2")
    agrupadas_colunas = (
        ("grupo", "Semana"), ("leituras", "Leituras"),
        ("armadilhas", "Armadilhas lidas"), ("positivas", "Leituras positivas"),
        ("ovos", "Ovos"), ("ipo", "IPO (%)"), ("ido", "IDO"), ("imo", "IMO"),
    )
    _write_table_sheet(wb, "Semanas", agrupadas_colunas, semanas)
    _write_table_sheet(
        wb, "Localidades",
        (("grupo", "Localidade"),) + agrupadas_colunas[1:],
        localidades,
    )
    ocorrencias_colunas = tuple(
        coluna for coluna in LEITURAS_COLUNAS
        if coluna[0] in {
            "id_contagem", "ovitrampa_id", "ano", "semana", "data", "ovos",
            "localidade", "rua", "numero", "complemento", "quarteirao",
            "lote_id", "diario_nome", "movimento", "ciclo", "lote_status",
            "laboratorista_nome", "ocorrencia_codigo", "ocorrencia",
        }
    )
    _write_table_sheet(wb, "Ocorrências", ocorrencias_colunas, ocorrencias, freeze="C2")

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = AZUL if ws.title == "Resumo" else CINZA_CLARO
    wb["Leituras"].sheet_properties.tabColor = AZUL_CLARO
    wb["Armadilhas"].sheet_properties.tabColor = VERDE_CLARO

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    nome = f"ovitrampas_monitoramento_{gerado_em.strftime('%Y%m%d_%H%M')}.xlsx"
    return output, nome
