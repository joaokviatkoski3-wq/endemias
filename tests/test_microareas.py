import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree as ET

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import load_workbook

from app_core import microareas
from app_core import registro_geografico as rg_core


def camada(quarteiroes=("0007", "0008"), longitude=-49.30):
    return json.dumps({"type": "FeatureCollection", "features": [
        {"type": "Feature", "properties": {"Localidade": "Sede", "id_Q": q},
         "geometry": {"type": "Polygon", "coordinates": [[
             [longitude + i * .01, -25.30], [longitude + i * .01 + .005, -25.30],
             [longitude + i * .01, -25.31], [longitude + i * .01, -25.30],
         ]]}}
        for i, q in enumerate(quarteiroes)
    ]}, ensure_ascii=False).encode("utf-8")


def desenho(pontos):
    return {"type": "Polygon", "coordinates": [[*pontos, pontos[0]]]}


LADO_1 = desenho([[-49.2999, -25.3001], [-49.298, -25.3001], [-49.2999, -25.303]])
LADO_2 = desenho([[-49.297, -25.3001], [-49.296, -25.3001], [-49.297, -25.302]])


def parte(lado, geometry):
    return {"quarteirao": "7", "logradouro": "Rua Teste", "lado": lado, "geometry": geometry}


class MicroareasTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base_dir = self.temp.name
        self.target = str(Path(self.base_dir) / "microareas.db")
        conn = sqlite3.connect(self.target)
        conn.row_factory = sqlite3.Row
        conn.executescript("""CREATE TABLE localidades (id_localidade INTEGER PRIMARY KEY, nome TEXT NOT NULL, cod_localidade TEXT);
            CREATE TABLE agentes (id_agente INTEGER PRIMARY KEY, nome TEXT, ativo INTEGER);
            INSERT INTO localidades VALUES (1, 'Sede', '11');
            INSERT INTO localidades VALUES (2, 'Graziela', '12');""")
        rg_core.ensure_schema(conn)
        conn.commit()
        conn.close()
        self.importar(camada())

    def importar(self, conteudo):
        previa = rg_core.preview_importacao_geojson(self.target, conteudo, "qgis.geojson", self.base_dir)
        rg_core.importar_geojson(self.target, conteudo, "qgis.geojson", previa["sha256"], self.base_dir)

    def preparar_lados(self):
        microareas.listar(self.target, self.base_dir)
        conn = sqlite3.connect(self.target)
        try:
            cursor = conn.execute("""INSERT INTO registro_geografico_quarteiroes
                (id_localidade, localidade, quarteirao, criado_em, atualizado_em)
                VALUES (1, 'Sede', '0007', '2026-09-29', '2026-09-29')""")
            for numero, lado, condominio in (("10", "1", 0), ("20", "2", 0), ("30", "3", 5)):
                conn.execute("""INSERT INTO registro_geografico_imoveis
                    (id_quarteirao, id_localidade, localidade, quarteirao, logradouro,
                     numero, lado, tipo, condominio, criado_em, atualizado_em)
                    VALUES (?, 1, 'Sede', '0007', 'Rua Teste', ?, ?, 'R', ?, '2026-09-29', '2026-09-29')""",
                    (cursor.lastrowid, numero, lado, condominio))
            conn.commit()
        finally:
            conn.close()

    def test_lados_parciais_entre_microareas_sem_contagem_duplicada(self):
        self.preparar_lados()
        inteira = microareas.salvar(self.target, {"id_localidade": 1, "numero": "1",
            "quarteiroes": ["7"]}, self.base_dir)
        microareas.salvar(self.target, {"id_localidade": 1, "numero": "1",
            "quarteiroes": [], "partes": [parte("1", LADO_1)]}, self.base_dir, inteira)
        segundo = microareas.salvar(self.target, {"id_localidade": 1, "numero": "2",
            "quarteiroes": [], "partes": [parte("2", LADO_2)]}, self.base_dir)
        dados = microareas.relatorio(self.target, self.base_dir)
        self.assertEqual(dados["indicadores"]["quarteiroes"], 1)
        self.assertEqual(dados["indicadores"]["lados_parciais"], 2)
        self.assertEqual(dados["indicadores"]["populacao_aproximada"], 6)
        self.assertEqual(microareas.relatorio(self.target, self.base_dir, incluir_condominios=False)["indicadores"]["populacao_aproximada"], 6)
        self.assertEqual(len(microareas.trechos(self.target, 1, "7", self.base_dir)["trechos"]), 3)
        self.assertIsNone(microareas.trechos(self.target, 1, "7", self.base_dir)["inteiro_id_microarea"])
        geo = microareas.exportar_geojson(self.target, self.base_dir)
        self.assertEqual(len(geo["features"]), 2)
        self.assertEqual({f["properties"]["lado"] for f in geo["features"]}, {"1", "2"})
        kml = ET.fromstring(microareas.exportar_kml(self.target, self.base_dir))
        self.assertEqual(len(kml.findall('.//{http://www.opengis.net/kml/2.2}Polygon')), 2)
        wb = load_workbook(io.BytesIO(microareas.exportar_xlsx(self.target, self.base_dir)))
        self.assertEqual(wb["Quarteirões"]["J2"].value, "Parcial")
        self.assertEqual(wb["Quarteirões"]["L3"].value, "2")
        with self.assertRaisesRegex(microareas.MicroareaError, "lados em outra"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "3",
                "quarteiroes": ["7"]}, self.base_dir)
        with self.assertRaisesRegex(microareas.MicroareaError, "já pertence"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "3",
                "quarteiroes": [], "partes": [parte("2", LADO_1)]}, self.base_dir)
        self.assertEqual(microareas.relatorio(self.target, self.base_dir, [segundo])["indicadores"]["populacao_aproximada"], 3)

    def test_desenhos_invalidos_sobrepostos_e_base_alterada(self):
        self.preparar_lados()
        um = microareas.salvar(self.target, {"id_localidade": 1, "numero": "1",
            "quarteiroes": [], "partes": [parte("1", LADO_1)]}, self.base_dir)
        with self.assertRaisesRegex(microareas.MicroareaError, "sobrepõe"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "2",
                "quarteiroes": [], "partes": [parte("2", LADO_1)]}, self.base_dir)
        with self.assertRaisesRegex(microareas.MicroareaError, "inteiro e parcial"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "2",
                "quarteiroes": ["7"], "partes": [parte("2", LADO_2)]}, self.base_dir)
        fora = desenho([[-49.31, -25.30], [-49.30, -25.30], [-49.31, -25.31]])
        dois = microareas.salvar(self.target, {"id_localidade": 1, "numero": "2",
            "quarteiroes": [], "partes": [parte("2", fora)]}, self.base_dir)
        salvo = microareas.listar(self.target, self.base_dir)["registros"]
        self.assertEqual(next(r for r in salvo if r["id_microarea"] == dois)["partes"][0]["geometry"]["coordinates"], fora["coordinates"])
        exportado = microareas.exportar_geojson(self.target, self.base_dir)
        self.assertEqual(next(f for f in exportado["features"] if f["properties"]["lado"] == "2")["geometry"]["coordinates"], fora["coordinates"])
        original = microareas.listar(self.target, self.base_dir)["registros"][0]["partes"][0]
        self.importar(camada(longitude=-49.31))
        self.assertTrue(microareas.listar(self.target, self.base_dir)["registros"][0]["partes"][0]["base_desatualizada"])
        with self.assertRaisesRegex(microareas.MicroareaError, "Redesenhe"):
            microareas.exportar_geojson(self.target, self.base_dir)
        with self.assertRaisesRegex(microareas.MicroareaError, "Redesenhe"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "1", "quarteiroes": [],
                "partes": [{**parte("1", LADO_1), "base_geometry_hash": original["base_geometry_hash"]}]},
                self.base_dir, um)

    def test_desenho_pode_cruzar_contorno_oficial(self):
        self.preparar_lados()
        cruzando = desenho([[-49.297, -25.301], [-49.294, -25.301], [-49.297, -25.302]])
        microareas.salvar(self.target, {"id_localidade": 1, "numero": "1",
            "quarteiroes": [], "partes": [parte("1", cruzando)]}, self.base_dir)
        self.assertEqual(microareas.exportar_geojson(self.target, self.base_dir)["features"][0]["geometry"]["coordinates"], cruzando["coordinates"])
        with self.assertRaisesRegex(microareas.MicroareaError, "sobrepõe"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "2",
                "quarteiroes": [], "partes": [parte("2", cruzando)]}, self.base_dir)

    def test_lado_removido_do_rg_exige_revisao(self):
        self.preparar_lados()
        microareas.salvar(self.target, {"id_localidade": 1, "numero": "1",
            "quarteiroes": [], "partes": [parte("1", LADO_1)]}, self.base_dir)
        conn = sqlite3.connect(self.target)
        try:
            conn.execute("UPDATE registro_geografico_imoveis SET lado='4' WHERE quarteirao='0007' AND lado='1'")
            conn.commit()
        finally:
            conn.close()
        self.assertTrue(microareas.listar(self.target, self.base_dir)["registros"][0]["partes"][0]["lado_ausente_rg"])
        with self.assertRaisesRegex(microareas.MicroareaError, "não consta mais no RG"):
            microareas.exportar_geojson(self.target, self.base_dir)

    def test_salvar_editar_e_impedir_sobreposicao(self):
        um = microareas.salvar(self.target, {"id_localidade": 1, "numero": "01", "quarteiroes": ["7"]}, self.base_dir)
        self.assertEqual(microareas.listar(self.target, self.base_dir)["registros"][0]["numero"], "1")
        with self.assertRaisesRegex(microareas.MicroareaError, "existe"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "1", "quarteiroes": ["8"]}, self.base_dir)
        with self.assertRaisesRegex(microareas.MicroareaError, "já pertence"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "2", "quarteiroes": ["7"]}, self.base_dir)
        microareas.salvar(self.target, {"id_localidade": 1, "numero": "1", "quarteiroes": ["7", "8"]}, self.base_dir, um)
        self.assertEqual(set(microareas.listar(self.target, self.base_dir)["registros"][0]["quarteiroes"]), {"0007", "0008"})
        with self.assertRaisesRegex(microareas.MicroareaError, "camada"):
            microareas.salvar(self.target, {"id_localidade": 2, "numero": "1", "quarteiroes": ["7"]}, self.base_dir)
        with self.assertRaisesRegex(microareas.MicroareaError, "catálogo"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "2", "acs_codigo": "acs-006", "quarteiroes": ["8"]}, self.base_dir)

    def test_atualizacao_geom_preserva_vinculos_e_exportacoes(self):
        microareas.listar(self.target, self.base_dir)
        conn = sqlite3.connect(self.target)
        try:
            conn.execute("INSERT INTO acs_catalogo VALUES ('acs-006', 'Maria', '2026-09-28')")
            conn.commit()
        finally:
            conn.close()
        identificador = microareas.salvar(self.target, {"id_localidade": 1, "numero": "1", "acs_codigo": "acs-006", "quarteiroes": ["7", "8"]}, self.base_dir)
        self.importar(camada(("0007",), -49.31))
        reg = microareas.listar(self.target, self.base_dir)["registros"][0]
        self.assertEqual(reg["acs_nome"], "Maria")
        self.assertEqual(reg["quarteiroes_sem_geometria"], ["0008"])
        self.assertEqual(microareas.exportar_geojson(self.target, self.base_dir)["features"][0]["geometry"]["coordinates"][0][0][0], -49.31)
        self.assertEqual(len(microareas.exportar_geojson(self.target, self.base_dir)["features"]), 1)
        kml = ET.fromstring(microareas.exportar_kml(self.target, self.base_dir))
        self.assertEqual(len(kml.findall('.//{http://www.opengis.net/kml/2.2}Placemark')), 1)
        wb = load_workbook(io.BytesIO(microareas.exportar_xlsx(self.target, self.base_dir)))
        self.assertEqual(wb["Microáreas"].max_row, 2)
        self.assertEqual(wb["Quarteirões"].max_row, 3)
        self.assertEqual(wb["Microáreas"]["F2"].value, 1)
        with self.assertRaisesRegex(microareas.MicroareaError, "camada"):
            microareas.salvar(self.target, {"id_localidade": 1, "numero": "1", "quarteiroes": ["7", "9"]}, self.base_dir, identificador)
        microareas.excluir(self.target, identificador)
        self.assertEqual(microareas.listar(self.target, self.base_dir)["registros"], [])
        microareas.salvar(self.target, {"id_localidade": 1, "numero": "1", "quarteiroes": ["7"]}, self.base_dir)

    def test_mesmo_numero_em_localidades_distintas(self):
        documento = json.loads(camada(("0007",)).decode("utf-8"))
        outro = json.loads(json.dumps(documento["features"][0]))
        outro["properties"]["Localidade"] = "Graziela"
        documento["features"].append(outro)
        self.importar(json.dumps(documento, ensure_ascii=False).encode("utf-8"))
        primeiro = microareas.salvar(self.target, {"id_localidade": 1, "numero": "1", "quarteiroes": ["7"]}, self.base_dir)
        segundo = microareas.salvar(self.target, {"id_localidade": 2, "numero": "1", "quarteiroes": ["7"]}, self.base_dir)
        self.assertNotEqual(primeiro, segundo)
        self.assertEqual(len(microareas.listar(self.target, self.base_dir)["registros"]), 2)

    def test_mesmo_acs_em_varias_microareas_e_troca_de_responsavel(self):
        microareas.listar(self.target, self.base_dir)
        conn = sqlite3.connect(self.target)
        try:
            conn.executemany("INSERT INTO acs_catalogo VALUES (?, ?, ?)", [
                ("acs-006", "Maria", "2026-09-29"),
                ("acs-007", "Ana", "2026-09-29"),
            ])
            conn.commit()
        finally:
            conn.close()
        primeiro = microareas.salvar(self.target, {
            "id_localidade": 1, "numero": "1", "acs_codigo": "acs-006", "quarteiroes": ["7"],
        }, self.base_dir)
        segundo = microareas.salvar(self.target, {
            "id_localidade": 1, "numero": "2", "acs_codigo": "acs-006", "quarteiroes": ["8"],
        }, self.base_dir)
        registros = {r["id_microarea"]: r for r in microareas.listar(self.target, self.base_dir)["registros"]}
        self.assertEqual(registros[primeiro]["acs_codigo"], "acs-006")
        self.assertEqual(registros[segundo]["acs_codigo"], "acs-006")
        microareas.salvar(self.target, {
            "id_localidade": 1, "numero": "2", "acs_codigo": "acs-007", "quarteiroes": ["8"],
        }, self.base_dir, segundo)
        registros = {r["id_microarea"]: r for r in microareas.listar(self.target, self.base_dir)["registros"]}
        self.assertEqual(registros[primeiro]["acs_nome"], "Maria")
        self.assertEqual(registros[segundo]["acs_nome"], "Ana")
        self.assertEqual(registros[segundo]["quarteiroes"], ["0008"])

    def test_relatorio_distingue_populacao_estimavel_de_rg_ausente(self):
        microareas.listar(self.target, self.base_dir)
        conn = sqlite3.connect(self.target)
        try:
            conn.executemany("INSERT INTO acs_catalogo VALUES (?, ?, ?)", [
                ("acs-006", "Maria", "2026-09-29"),
                ("acs-007", "Ana", "2026-09-29"),
            ])
            cursor = conn.execute("""INSERT INTO registro_geografico_quarteiroes
                (id_localidade, localidade, quarteirao, criado_em, atualizado_em)
                VALUES (1, 'Sede', '0007', '2026-09-29', '2026-09-29')""")
            conn.execute("""INSERT INTO registro_geografico_imoveis
                (id_quarteirao, id_localidade, localidade, quarteirao, logradouro, numero, tipo, criado_em, atualizado_em)
                VALUES (?, 1, 'Sede', '0007', 'Rua Teste', '10', 'R', '2026-09-29', '2026-09-29')""", (cursor.lastrowid,))
            conn.commit()
        finally:
            conn.close()
        primeiro = microareas.salvar(self.target, {
            "id_localidade": 1, "numero": "1", "acs_codigo": "acs-006", "quarteiroes": ["7"],
        }, self.base_dir)
        segundo = microareas.salvar(self.target, {
            "id_localidade": 1, "numero": "2", "quarteiroes": ["8"],
        }, self.base_dir)
        dados = microareas.relatorio(self.target, self.base_dir)
        self.assertEqual(dados["indicadores"]["microareas"], 2)
        self.assertEqual(dados["indicadores"]["sem_acs"], 1)
        self.assertEqual(dados["indicadores"]["acs_catalogo_sem_area_global"], 1)
        self.assertEqual(dados["indicadores"]["diferenca_cadastral_1a1"], 0)
        self.assertEqual(dados["indicadores"]["quarteiroes_sem_rg"], 1)
        self.assertEqual(dados["indicadores"]["populacao_aproximada"], 3)
        self.assertEqual(dados["indicadores"]["media_populacao_por_acs_com_rg"], 3.0)
        self.assertEqual(dados["por_acs"][0]["acs_codigo"], "acs-006")
        self.assertEqual(dados["por_localidade"][0]["sem_acs"], 1)
        recorte = microareas.relatorio(self.target, self.base_dir, [segundo])
        self.assertEqual(recorte["indicadores"]["populacao_aproximada"], 0)
        self.assertIsNone(recorte["indicadores"]["media_populacao_por_acs_com_rg"])
        self.assertEqual(recorte["indicadores"]["acs_catalogo_sem_area_global"], 1)
        wb = load_workbook(io.BytesIO(microareas.exportar_xlsx(self.target, self.base_dir)))
        self.assertEqual(wb.sheetnames, ["Indicadores", "Por ACS", "Por localidade", "Microáreas", "Quarteirões"])
        self.assertEqual(wb["Indicadores"]["B5"].value, 1)
        self.assertEqual(wb["Por ACS"]["F2"].value, 3)
        self.assertIsNone(wb["Microáreas"]["J3"].value)
        self.assertIsNone(wb["Quarteirões"]["H3"].value)
        ambiente = Environment(loader=FileSystemLoader(Path(__file__).resolve().parents[1] / "templates"),
                               autoescape=select_autoescape(["html"]))
        html = ambiente.get_template("microareas_relatorio.html").render(dados=dados)
        self.assertIn("Relatório de microáreas dos ACS", html)
        self.assertIn("Maria (acs-006)", html)
        self.assertIn("Quarteirões sem RG", html)

    def test_condominios_podem_ser_excluidos_sem_alterar_rg(self):
        microareas.listar(self.target, self.base_dir)
        conn = sqlite3.connect(self.target)
        try:
            cursor = conn.execute("""INSERT INTO registro_geografico_quarteiroes
                (id_localidade, localidade, quarteirao, criado_em, atualizado_em)
                VALUES (1, 'Sede', '0007', '2026-09-29', '2026-09-29')""")
            for numero, unidades in (("10", 0), ("20", 10)):
                conn.execute("""INSERT INTO registro_geografico_imoveis
                    (id_quarteirao, id_localidade, localidade, quarteirao, logradouro, numero, tipo, condominio, criado_em, atualizado_em)
                    VALUES (?, 1, 'Sede', '0007', 'Rua Teste', ?, 'R', ?, '2026-09-29', '2026-09-29')""",
                    (cursor.lastrowid, numero, unidades))
            conn.commit()
        finally:
            conn.close()
        microareas.salvar(self.target, {"id_localidade": 1, "numero": "1", "quarteiroes": ["7", "8"]}, self.base_dir)
        rg = rg_core.resumo_mapa(self.target, self.base_dir)["quarteiroes"][f"1:{rg_core._quarteirao_display('0007')}"]
        self.assertEqual(rg["residencias_condominio"], 10)
        self.assertEqual(rg["populacao_aproximada"], 32)
        self.assertEqual(rg["populacao_sem_condominios"], 3)
        com = microareas.relatorio(self.target, self.base_dir)
        sem = microareas.relatorio(self.target, self.base_dir, incluir_condominios=False)
        self.assertEqual(com["indicadores"]["populacao_aproximada"], 32)
        self.assertEqual(sem["indicadores"]["populacao_aproximada"], 3)
        self.assertEqual(sem["indicadores"]["quarteiroes_sem_rg"], 1)
        self.assertEqual(sem["registros"][0]["populacao_com_condominios"], 32)
        self.assertEqual(sem["registros"][0]["populacao_sem_condominios"], 3)
        wb = load_workbook(io.BytesIO(microareas.exportar_xlsx(
            self.target, self.base_dir, incluir_condominios=False)))
        self.assertEqual(wb["Indicadores"]["B2"].value, "Excluída")
        self.assertEqual(wb["Microáreas"]["J2"].value, 3)
        self.assertEqual(wb["Quarteirões"]["H2"].value, 3)
        self.assertIsNone(wb["Quarteirões"]["H3"].value)
        ambiente = Environment(loader=FileSystemLoader(Path(__file__).resolve().parents[1] / "templates"),
                               autoescape=select_autoescape(["html"]))
        html = ambiente.get_template("microareas_relatorio.html").render(dados=sem)
        self.assertIn("A4 portrait", html)
        self.assertIn("excluídos da estimativa", html)


if __name__ == "__main__":
    unittest.main()
