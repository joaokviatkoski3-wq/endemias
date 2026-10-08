import io
import sqlite3
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from flask import Flask
from openpyxl import load_workbook

from app_core import registro_geografico as rg
from app_core import registro_geografico_exportacao as exportacao
from blueprints.auth import bp as auth_bp
from blueprints.registro_geografico import bp


class ExportacaoRgTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.target = str(Path(self.tmp.name) / "rg.db")
        conn = sqlite3.connect(self.target)
        conn.row_factory = sqlite3.Row
        conn.executescript("""
            CREATE TABLE localidades (id_localidade INTEGER PRIMARY KEY, nome TEXT);
            CREATE TABLE agentes (id_agente INTEGER PRIMARY KEY, nome TEXT, ativo INTEGER);
            CREATE TABLE usuarios (id_usuario INTEGER PRIMARY KEY, ativo INTEGER, nivel TEXT);
            INSERT INTO localidades VALUES (1, 'São Venâncio'), (2, 'Sede');
            INSERT INTO agentes VALUES (1, 'João', 1), (2, 'Maria', 1);
            INSERT INTO usuarios VALUES (1, 1, 'visualizador'), (2, 0, 'admin');
        """)
        rg.ensure_schema(conn)
        for localidade, numero in ((1, "0001"), (1, "0408.1"), (1, "0003"), (2, "0001")):
            conn.execute("""INSERT INTO registro_geografico_quarteiroes
                (id_localidade, localidade, quarteirao, criado_em, atualizado_em)
                VALUES (?, ?, ?, '2026-10-08', '2026-10-08')""",
                (localidade, "São Venâncio" if localidade == 1 else "Sede", numero))
        for ordem, numero, seq, tipo, cond, data, obs in (
            (2, "SN", None, "REF", None, None, "Referência do lado"),
            (1, "0012A", "02", "R", 20, "2026-10-08", '=HYPERLINK("https://example.org")'),
        ):
            conn.execute("""INSERT INTO registro_geografico_imoveis
                (id_quarteirao, ordem, id_localidade, localidade, quarteirao, logradouro,
                 numero, sequencia, lado, tipo, condominio, data_atualizacao, observacao,
                 agentes_texto, criado_em, atualizado_em)
                VALUES (1, ?, 1, 'São Venâncio', '0001', 'Rua São João', ?, ?, '01', ?, ?, ?, ?,
                        'Agente legado', '2026-10-08', '2026-10-08')""",
                (ordem, numero, seq, tipo, cond, data, obs))
        conn.executescript("""
            INSERT INTO registro_geografico_imovel_agentes VALUES (2, 1), (2, 2);
            INSERT INTO registro_geografico_imoveis
                (id_quarteirao, ordem, id_localidade, localidade, quarteirao, logradouro,
                 numero, tipo, condominio, criado_em, atualizado_em)
                VALUES (2, 1, 1, 'São Venâncio', '0408.1', 'Rua B', '0', 'C', 0, '2026-10-08', '2026-10-08'),
                       (4, 1, 2, 'Sede', '0001', 'Não exportar', '9', 'O', 0, '2026-10-08', '2026-10-08');
        """)
        conn.commit()
        conn.close()

    def workbook(self, quarteiroes, localidade=1):
        wb = load_workbook(io.BytesIO(exportacao.gerar_xlsx(self.target, localidade, quarteiroes)))
        self.addCleanup(wb.close)
        return wb

    def test_unico_rg_com_dados_completos_e_sem_metadados(self):
        wb = self.workbook(["1"])
        self.assertEqual(wb.sheetnames, ["RGs"])
        ws = wb.active
        self.assertEqual(tuple(c.value for c in ws[1]), exportacao.COLUNAS)
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.max_column, 11)
        self.assertEqual([ws[f"D{r}"].value for r in (2, 3)], ["0012A", "SN"])
        self.assertEqual(ws["E2"].value, "02")
        self.assertEqual(ws["F2"].value, "01")
        self.assertEqual(ws["H2"].value, 20)
        self.assertIsInstance(ws["H2"].value, int)
        self.assertEqual(ws["J2"].value, datetime(2026, 10, 8))
        self.assertEqual(ws["J2"].number_format, "dd/mm/yyyy")
        self.assertEqual(set(ws["K2"].value.split(", ")), {"João", "Maria"})
        self.assertEqual(ws["K3"].value, "Agente legado")
        self.assertEqual(ws["G3"].value, "REF")
        self.assertIsNone(ws["J3"].value)
        self.assertEqual(ws["I2"].value, '=HYPERLINK("https://example.org")')
        self.assertEqual(ws["I2"].data_type, "s")
        self.assertFalse(any(c.data_type == "f" for row in ws for c in row))
        self.assertFalse(ws.merged_cells.ranges)
        self.assertEqual(ws.freeze_panes, "C2")
        self.assertEqual(ws.auto_filter.ref, "A1:K3")

    def test_multiplos_sem_duplicar_e_sem_outros_rgs_ou_localidades(self):
        ws = self.workbook(["0408.1", "1", "0001", "0408.1"]).active
        self.assertEqual(ws.max_row, 4)
        self.assertEqual([ws[f"B{r}"].value for r in range(2, 5)], ["0408.1", "1", "1"])
        self.assertEqual(ws["B2"].data_type, "s")
        self.assertEqual(ws["D2"].value, "0")
        self.assertEqual(ws["H2"].value, 0)
        self.assertTrue(all(ws[f"A{r}"].value == "São Venâncio" for r in range(2, 5)))
        self.assertNotIn("Não exportar", str(list(ws.values)))

    def test_rg_sem_imoveis_preserva_identificacao_sem_inventar_dados(self):
        ws = self.workbook(["3"]).active
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws["A2"].value, "São Venâncio")
        self.assertEqual(ws["B2"].value, "3")
        self.assertTrue(all(c.value is None for c in ws[2][2:]))

    def test_selecao_obrigatoria_e_quarteiroes_da_localidade(self):
        for loc, qs in ((None, ["1"]), ("abc", ["1"]), (0, ["1"]),
                        (1, []), (1, [" "]), (1, ["1", "9999"]), (2, ["0408.1"]), (99, ["1"])):
            with self.subTest(loc=loc, qs=qs), self.assertRaises(ValueError):
                exportacao.gerar_xlsx(self.target, loc, qs)

    def test_nao_limita_exportacao_a_paginacao_da_consulta(self):
        conn = sqlite3.connect(self.target)
        conn.executemany("""INSERT INTO registro_geografico_imoveis
            (id_quarteirao, ordem, id_localidade, localidade, quarteirao, logradouro,
             numero, tipo, criado_em, atualizado_em)
            VALUES (1, ?, 1, 'São Venâncio', '0001', 'Rua Grande', ?, 'R',
                    '2026-10-08', '2026-10-08')""",
            [(i + 3, str(i + 1)) for i in range(650)])
        conn.commit()
        conn.close()
        ws = self.workbook(["1"]).active
        self.assertEqual(ws.max_row, 653)
        self.assertEqual(ws["D653"].value, "650")

    def test_http_autenticacao_download_e_erros(self):
        app = Flask(__name__)
        app.secret_key = "testes"
        app.config.update(TESTING=True, DB_BACKEND="sqlite", DB_PATH=self.target)
        app.register_blueprint(auth_bp)
        app.register_blueprint(bp)
        client = app.test_client()
        url = "/registro-geografico/exportar.xlsx?localidade=1&quarteirao=1&quarteirao=0408.1"
        self.assertEqual(client.get(url).status_code, 302)
        with client.session_transaction() as sess:
            sess["uid"] = 1
        resposta = client.get(url)
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment; filename=rgs_selecionados.xlsx", resposta.headers["Content-Disposition"])
        wb = load_workbook(io.BytesIO(resposta.data))
        self.assertEqual(wb.active.max_row, 4)
        wb.close()
        self.assertEqual(client.get("/registro-geografico/exportar.xlsx?localidade=1").status_code, 400)
        resposta = client.get("/registro-geografico/exportar.xlsx?localidade=abc&quarteirao=1")
        self.assertEqual(resposta.status_code, 400)
        self.assertIn("localidade", resposta.json["erro"])
        with client.session_transaction() as sess:
            sess["uid"] = 2
        self.assertEqual(client.get(url).status_code, 302)


if __name__ == "__main__":
    unittest.main()
