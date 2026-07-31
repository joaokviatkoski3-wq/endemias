"""Homologa o Boletim Mensal no PostgreSQL sem alterar dados publicos."""

import argparse
import io
import json
import logging
import os
import sys
import tempfile
import traceback
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app_core import db as db_core  # noqa: E402


SAFE_DATABASE = "endemias_teste"
TEMP_TABLES = (
    "usuarios",
    "auditoria_eventos",
    "boletim_mensal_itens",
    "visitas",
    "depositos_inspecionados",
    "coletas",
    "focos_positivos",
    "esporotricose_visitas",
    "esporotricose_animais",
    "recolhimentos",
    "bri_registros",
    "amostras_animais",
    "acoes_setor",
)
ADMIN = {
    "id_usuario": 940001,
    "nome": "Admin Boletim PG",
    "nivel": "admin",
}
OPERADOR = {
    "id_usuario": 940002,
    "nome": "Operador Boletim PG",
    "nivel": "operador",
}
VISUALIZADOR = {
    "id_usuario": 940003,
    "nome": "Visualizador Boletim PG",
    "nivel": "visualizador",
}


class _SharedConnection:
    """Mantem as tabelas temporarias disponiveis entre as rotas Flask."""

    def __init__(self, conn):
        self._conn = conn
        self.backend = conn.backend

    def close(self):
        pass

    def __enter__(self):
        self._conn.__enter__()
        return self

    def __exit__(self, exc_type, exc_value, traceback_obj):
        return self._conn.__exit__(exc_type, exc_value, traceback_obj)

    def __getattr__(self, name):
        return getattr(self._conn, name)


def _parser():
    parser = argparse.ArgumentParser(
        description=(
            "Testa o Boletim Mensal somente em tabelas PostgreSQL "
            "temporarias."
        )
    )
    parser.add_argument("--database", default=SAFE_DATABASE)
    parser.add_argument(
        "--confirmar-banco",
        help="Obrigatorio para qualquer banco diferente de endemias_teste.",
    )
    return parser


def _public_counts(conn):
    tables = [
        row[0]
        for row in conn.execute(
            """SELECT table_name
                 FROM information_schema.tables
                WHERE table_schema='public'
                  AND table_type='BASE TABLE'
                ORDER BY table_name"""
        ).fetchall()
    ]
    return {
        table: conn.execute(
            f"SELECT COUNT(*) FROM public.{table}"
        ).fetchone()[0]
        for table in tables
    }


def _temporary_schema(conn):
    for table in TEMP_TABLES:
        conn.execute(
            f"""CREATE TEMPORARY TABLE {table}
                (LIKE public.{table} INCLUDING ALL)
                ON COMMIT PRESERVE ROWS"""
        )
    conn.commit()


def _fixtures(conn):
    for usuario, login in (
        (ADMIN, "admin_boletim_pg"),
        (OPERADOR, "operador_boletim_pg"),
        (VISUALIZADOR, "visualizador_boletim_pg"),
    ):
        conn.execute(
            """INSERT INTO usuarios
               (id_usuario, usuario, nome, senha_hash, nivel, ativo, criado_em,
                acesso_laboratorio, somente_laboratorio)
               VALUES (?, ?, ?, 'teste', ?, 1, '2026-07-31T08:00:00', 0, 0)""",
            (
                usuario["id_usuario"],
                login,
                usuario["nome"],
                usuario["nivel"],
            ),
        )

    for tipo in ("PVE", "TB", "TBO", "PE"):
        conn.execute(
            """INSERT INTO visitas
               (id_visita, kobo_uuid, tipo, data, processado_em)
               VALUES (?, ?, ?, '2026-07-15', '2026-07-15T12:00:00')""",
            (f"bm-visita-{tipo.lower()}", f"bm-kobo-{tipo.lower()}", tipo),
        )
    conn.execute(
        """INSERT INTO depositos_inspecionados
           (id_visita, tipo_deposito, inspecionado, eliminado, tratado)
           VALUES ('bm-visita-pve', 'A1', 3, 2, 1)"""
    )
    conn.execute(
        """INSERT INTO coletas
           (id_coleta, id_visita, num_tubo)
           VALUES ('bm-coleta-1', 'bm-visita-pve', 'BM-001')"""
    )
    conn.execute(
        """INSERT INTO focos_positivos
           (id_foco, id_visita, id_coleta, data, origem, gera_notificacao)
           VALUES ('bm-foco-1', 'bm-visita-pve', 'bm-coleta-1',
                   '2026-07-15', 'sistema', 1)"""
    )

    conn.execute(
        """INSERT INTO esporotricose_visitas
           (id_visita, kobo_uuid, data, processado_em)
           VALUES ('bm-esporo-visita-1', 'bm-esporo-kobo-1', '2026-07-16',
                   '2026-07-16T12:00:00')"""
    )
    conn.execute(
        """INSERT INTO esporotricose_animais
           (id_animal, id_visita, feridas, processado_em)
           VALUES
             ('bm-animal-1', 'bm-esporo-visita-1', 'sim',
              '2026-07-16T12:00:00'),
             ('bm-animal-2', 'bm-esporo-visita-1', 'nao',
              '2026-07-16T12:00:00')"""
    )
    conn.execute(
        """INSERT INTO recolhimentos
           (id_recolhimento, data, pneu, total_materiais, origem_estrutura,
            processado_em)
           VALUES ('bm-recolhimento-1', '2026-07-17', 4, 6, 'nova',
                   '2026-07-17T12:00:00')"""
    )
    conn.execute(
        """INSERT INTO bri_registros
           (id_bri, data, origem_estrutura, processado_em)
           VALUES ('bm-bri-1', '2026-07-18', 'nova',
                   '2026-07-18T12:00:00')"""
    )
    conn.execute(
        """INSERT INTO amostras_animais
           (id_amostra, data, quantidade, origem_estrutura, processado_em)
           VALUES ('bm-amostra-1', '2026-07-19', 3, 'nova',
                   '2026-07-19T12:00:00')"""
    )
    for tipo, dia in (
        ("educativa", "20"),
        ("limpeza", "21"),
        ("reuniao", "22"),
    ):
        conn.execute(
            """INSERT INTO acoes_setor
               (tipo, situacao, data, criado_em)
               VALUES (?, 'realizada', ?, '2026-07-22T12:00:00')""",
            (tipo, f"2026-07-{dia}"),
        )
    conn.commit()


def _login(client, usuario):
    with client.session_transaction() as flask_session:
        flask_session["uid"] = usuario["id_usuario"]
        flask_session["nome"] = usuario["nome"]
        flask_session["nivel"] = usuario["nivel"]


def _assert_status(response, expected, label):
    if response.status_code != expected:
        raise RuntimeError(
            f"{label} respondeu HTTP {response.status_code}: "
            f"{response.get_data(as_text=True)}"
        )
    return response


def _linha(dados, chave):
    try:
        return next(item for item in dados["linhas"] if item["chave"] == chave)
    except StopIteration as exc:
        raise RuntimeError(f"Indicador ausente no boletim: {chave}") from exc


def _validar_automaticos(dados):
    esperados = {
        "visitas_pve": 1,
        "visitas_tb": 1,
        "visitas_tbo": 1,
        "visitas_pe": 1,
        "depositos_inspecionados": 3,
        "depositos_eliminados": 2,
        "coletas_laboratorio": 1,
        "focos_positivos": 1,
        "esporotricose_visitas": 1,
        "esporotricose_animais": 2,
        "esporotricose_animais_feridas": 1,
        "recolhimentos_materiais": 6,
        "recolhimentos_pneus": 4,
        "bri_registros": 1,
        "amostras_animais_registros": 1,
        "amostras_animais_quantidade": 3,
        "acoes_setor_total": 2,
    }
    encontrados = {
        chave: _linha(dados, chave)["quantidade"]
        for chave in esperados
    }
    if encontrados != esperados:
        raise RuntimeError(
            f"Indicadores automaticos divergiram: {encontrados!r}"
        )
    if dados["total"] != 32:
        raise RuntimeError(
            f"Total automatico esperado 32, encontrado {dados['total']}."
        )


def _payload_fechamento():
    return {
        "mes": "2026-07",
        "linhas": [
            {
                "chave": "visitas_pve",
                "origem": "auto",
                "ordem": 10,
                "indicador": "Vistorias PVE ajustadas no fechamento",
                "quantidade": 9,
                "unidade": "visitas",
                "ativo": True,
            },
            {
                "chave": "manual_educacao_pg",
                "origem": "manual",
                "ordem": 180,
                "indicador": "Atividade educativa complementar PG",
                "quantidade": 5,
                "unidade": "atividades",
                "ativo": True,
            },
            {
                "chave": "manual_inativo_pg",
                "origem": "manual",
                "ordem": 190,
                "indicador": "Linha manual inativa PG",
                "quantidade": 99,
                "unidade": "registros",
                "ativo": False,
            },
        ],
    }


def _validar_xlsx(response):
    workbook = openpyxl.load_workbook(
        io.BytesIO(response.data),
        read_only=True,
        data_only=True,
    )
    try:
        values = {
            cell.value
            for row in workbook.active.iter_rows()
            for cell in row
            if cell.value is not None
        }
    finally:
        workbook.close()
    if "Atividade educativa complementar PG" not in values:
        raise RuntimeError("A linha manual nao apareceu na exportacao XLSX.")
    if 45 not in values:
        raise RuntimeError("O total ajustado nao apareceu na exportacao XLSX.")


def _test_routes(database, conn, create_app):
    with tempfile.TemporaryDirectory(prefix="endemias-pg-boletim-") as tmpdir:
        temp_path = Path(tmpdir)
        log_path = str(temp_path / "teste.log")
        try:
            flask_app = create_app(
                {
                    "DB_BACKEND": "postgresql",
                    "PG_DATABASE": database,
                    "TESTING": True,
                    "WTF_CSRF_ENABLED": False,
                    "LOG_PATH": log_path,
                    "SECRET_KEY_PATH": str(temp_path / "secret.key"),
                    "ANEXOS_DIR": str(temp_path / "anexos"),
                    "BACKUP_DIR": str(temp_path / "backups"),
                }
            )
            admin_client = flask_app.test_client()
            operador_client = flask_app.test_client()
            visualizador_client = flask_app.test_client()
            _login(admin_client, ADMIN)
            _login(operador_client, OPERADOR)
            _login(visualizador_client, VISUALIZADOR)

            _assert_status(
                visualizador_client.get("/boletim-mensal"),
                200,
                "Pagina do Boletim Mensal",
            )
            automaticos = _assert_status(
                admin_client.get(
                    "/api/boletim-mensal?mes=2026-07&modo=auto"
                ),
                200,
                "Indicadores automaticos",
            ).get_json()
            _validar_automaticos(automaticos)

            _assert_status(
                visualizador_client.post(
                    "/api/boletim-mensal",
                    json=_payload_fechamento(),
                ),
                403,
                "Bloqueio de edicao para visualizador",
            )
            salvo = _assert_status(
                operador_client.post(
                    "/api/boletim-mensal",
                    json=_payload_fechamento(),
                ),
                200,
                "Persistencia do fechamento mensal",
            ).get_json()
            if salvo["total"] != 45:
                raise RuntimeError(
                    f"Total ajustado esperado 45, encontrado {salvo['total']}."
                )
            if _linha(salvo, "visitas_pve")["quantidade"] != 9:
                raise RuntimeError("O ajuste do indicador automatico nao persistiu.")
            if _linha(salvo, "manual_educacao_pg")["quantidade"] != 5:
                raise RuntimeError("A linha manual nao persistiu.")
            if _linha(salvo, "manual_inativo_pg")["ativo"]:
                raise RuntimeError("A desativacao da linha manual nao persistiu.")

            recarregado = _assert_status(
                admin_client.get("/api/boletim-mensal?mes=2026-07"),
                200,
                "Consulta do fechamento salvo",
            ).get_json()
            if recarregado["total"] != 45:
                raise RuntimeError("O fechamento recarregado perdeu os ajustes.")

            automaticos_apos = _assert_status(
                admin_client.get(
                    "/api/boletim-mensal?mes=2026-07&modo=auto"
                ),
                200,
                "Recalculo automatico",
            ).get_json()
            _validar_automaticos(automaticos_apos)

            pdf = _assert_status(
                admin_client.get("/boletim-mensal/pdf?mes=2026-07"),
                200,
                "Relatorio para PDF",
            ).get_data(as_text=True)
            if "Atividade educativa complementar PG" not in pdf:
                raise RuntimeError("A linha manual nao apareceu no relatorio PDF.")

            xlsx = _assert_status(
                admin_client.get(
                    "/api/boletim-mensal/exportar?mes=2026-07"
                ),
                200,
                "Exportacao XLSX",
            )
            try:
                _validar_xlsx(xlsx)
            finally:
                xlsx.close()

            if conn.execute(
                """SELECT COUNT(*) FROM boletim_mensal_itens
                    WHERE ano_mes='2026-07'"""
            ).fetchone()[0] != 3:
                raise RuntimeError("A tabela temporaria nao recebeu os tres itens.")

            evento = conn.execute(
                """SELECT entidade_id, detalhes_json
                     FROM auditoria_eventos
                    WHERE acao='boletim_mensal_salvou'
                    ORDER BY id_evento DESC
                    LIMIT 1"""
            ).fetchone()
            if not evento or evento[0] != "2026-07":
                raise RuntimeError("A auditoria do fechamento nao foi registrada.")
            detalhes = json.loads(evento[1])
            if detalhes != {"ativos": 2, "itens": 3}:
                raise RuntimeError(
                    f"Detalhes da auditoria divergiram: {detalhes!r}"
                )
        finally:
            for handler in list(logging.getLogger().handlers):
                if getattr(handler, "baseFilename", None) == os.path.abspath(
                    log_path
                ):
                    logging.getLogger().removeHandler(handler)
                    handler.close()


def _test_temporary_data(database, create_app):
    target = db_core.DatabaseTarget("postgresql", database)
    conn = db_core.connect(target)
    original_connect = db_core.connect
    try:
        public_before = _public_counts(conn)
        conn.rollback()
        _temporary_schema(conn)
        _fixtures(conn)
        shared = _SharedConnection(conn)
        db_core.connect = lambda unused_target: shared

        _test_routes(database, conn, create_app)

        public_after = _public_counts(conn)
        if public_before != public_after:
            raise RuntimeError("Uma tabela publica foi alterada pelo ensaio.")
        return len(public_before)
    finally:
        db_core.connect = original_connect
        conn.close()


def main(argv=None):
    args = _parser().parse_args(argv)
    if (
        args.database != SAFE_DATABASE
        and args.confirmar_banco != args.database
    ):
        print(
            "[ERRO] Para testar outro banco, informe "
            f"--confirmar-banco {args.database}"
        )
        return 2

    try:
        from app import create_app

        total_publicas = _test_temporary_data(args.database, create_app)
    except Exception as exc:
        print(f"[ERRO] {exc}")
        traceback.print_exc()
        return 1

    print("Teste do Boletim Mensal no PostgreSQL")
    print("=" * 43)
    print(f"Banco: {args.database}")
    print("Indicadores automaticos e fontes operacionais: OK")
    print("Fechamento mensal, ajustes e linhas manuais: OK")
    print("Pagina, PDF e exportacao XLSX: OK")
    print("Permissoes e auditoria: OK")
    print(f"Tabelas publicas preservadas: {total_publicas}")
    print("\n[OK] Modulo validado somente em tabelas temporarias.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
