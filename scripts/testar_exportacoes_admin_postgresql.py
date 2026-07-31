"""Homologa Exportacoes, Consolidados e Central usando apenas tabelas temporarias."""

import argparse
import io
import logging
import os
import sys
import tempfile
import traceback
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app_core import db as db_core  # noqa: E402


SAFE_DATABASE = "endemias_teste"
ADMIN = {"id_usuario": 960001, "nome": "Admin Exportacoes PG", "nivel": "admin"}
ID_AGENTE = 960001
ID_LOCALIDADE = 960001


class _SharedConnection:
    def __init__(self, conn):
        self._conn = conn
        self.backend = conn.backend

    def close(self):
        pass

    def __getattr__(self, name):
        return getattr(self._conn, name)


def _parser():
    parser = argparse.ArgumentParser(
        description="Testa Exportacoes, Consolidados e Central no PostgreSQL."
    )
    parser.add_argument("--database", default=SAFE_DATABASE)
    parser.add_argument(
        "--confirmar-banco",
        help="Obrigatorio para qualquer banco diferente de endemias_teste.",
    )
    return parser


def _public_tables(conn):
    return [
        row[0]
        for row in conn.execute(
            """SELECT table_name FROM information_schema.tables
                WHERE table_schema='public' AND table_type='BASE TABLE'
                ORDER BY table_name"""
        ).fetchall()
    ]


def _public_counts(conn, tables):
    return {
        table: conn.execute(f'SELECT COUNT(*) FROM public."{table}"').fetchone()[0]
        for table in tables
    }


def _temporary_schema(conn, tables):
    for table in tables:
        conn.execute(
            f'''CREATE TEMPORARY TABLE "{table}"
                (LIKE public."{table}" INCLUDING ALL) ON COMMIT PRESERVE ROWS'''
        )
    conn.commit()


def _fixtures(conn):
    conn.execute(
        """INSERT INTO usuarios
           (id_usuario, usuario, nome, senha_hash, nivel, ativo, criado_em,
            acesso_laboratorio, somente_laboratorio)
           VALUES (?, 'admin_exportacoes_pg', ?, 'teste', 'admin', 1,
                   '2026-07-31T08:00:00', 1, 0)""",
        (ADMIN["id_usuario"], ADMIN["nome"]),
    )
    conn.execute(
        """INSERT INTO agentes(id_agente, nome, nome_completo, ativo)
           VALUES (?, 'agente_export_pg', 'Agente Exportacao PostgreSQL', 1)""",
        (ID_AGENTE,),
    )
    conn.execute(
        """INSERT INTO localidades(id_localidade, nome, cod_localidade)
           VALUES (?, 'Localidade Exportacao PG', 'LEX')""",
        (ID_LOCALIDADE,),
    )
    conn.execute(
        """INSERT INTO visitas
           (id_visita, kobo_uuid, tipo, data, hora_inicio, hora_fim,
            localidade, id_localidade, logradouro, numero, quarteirao,
            visita, morador, processado_em)
           VALUES ('exp-visita-1', 'exp-kobo-1', 'PE', '2026-07-28',
                   '08:00', '08:30', 'Localidade Exportacao PG', ?,
                   'Rua PostgreSQL', '10', 12, 'normal', 'Morador PG',
                   '2026-07-28T09:00:00')""",
        (ID_LOCALIDADE,),
    )
    conn.execute(
        "INSERT INTO visita_agentes(id_visita, id_agente) VALUES ('exp-visita-1', ?)",
        (ID_AGENTE,),
    )
    conn.execute(
        """INSERT INTO depositos_inspecionados
           (id_visita, tipo_deposito, inspecionado, eliminado, tratado,
            tipo_tratamento, qtd_carga)
           VALUES ('exp-visita-1', 'A1', 3, 1, 1, 'larvicida', 1)"""
    )
    conn.execute(
        """INSERT INTO tratamentos
           (id_visita, tipo, quantidade_carga, qtd_depositos_tratados)
           VALUES ('exp-visita-1', 'larvicida', 1, 2)"""
    )
    conn.execute(
        """INSERT INTO coletas
           (id_coleta, id_visita, num_tubo, codigo_deposito, tipo_deposito,
            deposito_eliminado)
           VALUES ('exp-coleta-1', 'exp-visita-1', 'PG-EXP-001', 'A1', 'pneu', 0)"""
    )
    conn.execute(
        """INSERT INTO resultados_laboratorio
           (id_coleta, num_tubo, data_coleta, laboratorista, data_leitura,
            aegypt_larvas, id_laboratorista, origem)
           VALUES ('exp-coleta-1', 'PG-EXP-001', '2026-07-28',
                   'Agente Exportacao PostgreSQL', '2026-07-29', 2, ?, 'sistema')""",
        (ID_AGENTE,),
    )
    conn.execute(
        """INSERT INTO focos_positivos
           (id_foco, id_visita, id_coleta, num_tubo, origem, tipo_trabalho,
            data, id_localidade, localidade, quarteirao, logradouro, numero,
            nome_morador, agentes, gera_notificacao, status_notificacao, codigo)
           VALUES ('exp-foco-1', 'exp-visita-1', 'exp-coleta-1', 'PG-EXP-001',
                   'sistema', 'PE', '2026-07-28', ?, 'Localidade Exportacao PG',
                   12, 'Rua PostgreSQL', '10', 'Morador PG',
                   'Agente Exportacao PostgreSQL', 1, 'pendente', 'EXP-PG-001')""",
        (ID_LOCALIDADE,),
    )
    conn.commit()


def _login(client):
    with client.session_transaction() as flask_session:
        flask_session.update({
            "uid": ADMIN["id_usuario"],
            "nome": ADMIN["nome"],
            "nivel": ADMIN["nivel"],
        })


def _response(client, path, method="get", **kwargs):
    response = getattr(client, method)(path, **kwargs)
    if response.status_code != 200:
        raise RuntimeError(
            f"{path} respondeu HTTP {response.status_code}: "
            f"{response.get_data(as_text=True)}"
        )
    return response


def _xlsx_rows(response):
    workbook = openpyxl.load_workbook(io.BytesIO(response.get_data()), read_only=True)
    try:
        return list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()


def _test_routes(database, create_app):
    with tempfile.TemporaryDirectory(prefix="endemias-pg-exportacoes-") as tmpdir:
        base = Path(tmpdir)
        log_path = str(base / "teste.log")
        app = create_app({
            "DB_BACKEND": "postgresql",
            "PG_DATABASE": database,
            "TESTING": True,
            "WTF_CSRF_ENABLED": False,
            "LOG_PATH": log_path,
            "SECRET_KEY_PATH": str(base / "secret.key"),
            "INSTANCE_DIR": str(base),
            "ANEXOS_DIR": str(base / "anexos"),
            "BACKUP_DIR": str(base / "backups"),
            "BACKUP_COMPLETO_DIR": str(base / "backups_completos"),
            "SAIDA_DIR": str(base / "saida"),
        })
        client = app.test_client()
        _login(client)
        try:
            visitas = _xlsx_rows(_response(
                client,
                "/api/visitas/exportar?d_ini=2026-07-01&d_fim=2026-07-31",
            ))
            if visitas[1][16] != "Agente Exportacao PostgreSQL":
                raise RuntimeError("A agregacao de agentes da exportacao divergiu.")

            notificacoes = _xlsx_rows(_response(
                client,
                "/api/notificacoes/exportar?busca=postgresql",
            ))
            if notificacoes[1][0] != "EXP-PG-001":
                raise RuntimeError("A busca textual da exportacao divergiu.")

            laboratorio = _xlsx_rows(_response(
                client,
                "/api/laboratorio/exportar?d_ini=2026-07-01&d_fim=2026-07-31&tubo=pg-exp",
            ))
            if laboratorio[1][-1] != "Agente Exportacao PostgreSQL":
                raise RuntimeError("A exportacao do laboratorio divergiu.")

            consolidados = _response(
                client,
                "/saida/gerar-consolidados",
                method="post",
                json={"tipo": "PE"},
            ).get_json()
            resultado = consolidados["resultados"][0]
            if resultado["visitas"] != 1 or resultado["coletas"] != 1:
                raise RuntimeError("Os totais do consolidado PostgreSQL divergiram.")

            central = _response(client, "/admin/sistema").get_data(as_text=True)
            if "O banco ativo é PostgreSQL" not in central:
                raise RuntimeError("A Central nao informou as limitacoes PostgreSQL.")
            diagnostico = _response(
                client,
                "/api/admin/sistema/diagnostico?completo=1",
            ).get_json()
            if diagnostico["resumo"]["backend"] != "postgresql":
                raise RuntimeError("O diagnostico nao identificou o PostgreSQL.")
            if not any(
                item["titulo"] == "Conexão PostgreSQL confirmada."
                for item in diagnostico["itens"]
            ):
                raise RuntimeError("A saude da conexao PostgreSQL nao foi confirmada.")

            backup_response = client.post("/admin/sistema/backups/criar")
            if backup_response.status_code != 302:
                raise RuntimeError("A Central nao concluiu o backup PostgreSQL.")
            dumps = list((base / "backups").glob("endemias_*.dump"))
            if len(dumps) != 1 or not dumps[0].with_suffix(".dump.json").is_file():
                raise RuntimeError("O dump PostgreSQL validado nao foi publicado.")
        except Exception:
            if os.path.exists(log_path):
                print(Path(log_path).read_text(encoding="utf-8", errors="replace"))
            raise
        finally:
            for handler in list(logging.getLogger().handlers):
                if getattr(handler, "baseFilename", None) == os.path.abspath(log_path):
                    logging.getLogger().removeHandler(handler)
                    handler.close()


def _test_temporary_data(database, create_app):
    target = db_core.DatabaseTarget("postgresql", database)
    conn = db_core.connect(target)
    original_connect = db_core.connect
    try:
        tables = _public_tables(conn)
        public_before = _public_counts(conn, tables)
        conn.rollback()
        _temporary_schema(conn, tables)
        _fixtures(conn)
        shared = _SharedConnection(conn)
        db_core.connect = lambda unused_target: shared
        _test_routes(database, create_app)
        conn.rollback()
        if public_before != _public_counts(conn, tables):
            raise RuntimeError("Uma tabela publica foi alterada pelo ensaio.")
        return len(tables)
    finally:
        db_core.connect = original_connect
        conn.close()


def main(argv=None):
    args = _parser().parse_args(argv)
    if args.database != SAFE_DATABASE and args.confirmar_banco != args.database:
        print(
            "[ERRO] Para testar outro banco, informe "
            f"--confirmar-banco {args.database}"
        )
        return 2
    try:
        from app import create_app

        total = _test_temporary_data(args.database, create_app)
    except Exception as exc:
        print(f"[ERRO] {exc}")
        traceback.print_exc()
        return 1

    print("Teste de Exportacoes, Consolidados e Central no PostgreSQL")
    print("=" * 62)
    print(f"Banco: {args.database}")
    print("Tres exportacoes XLSX e valores agregados: OK")
    print("Consolidado PE com abas e totais: OK")
    print("Central, diagnostico e backup PostgreSQL validado: OK")
    print(f"Tabelas publicas preservadas: {total}")
    print("\n[OK] Lote validado somente em tabelas temporarias.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
